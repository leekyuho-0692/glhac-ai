"""2026-10 감사 중간 결함 묶음 MB1~MB12 부정 테스트.

각 테스트는 실제 함수/엔드포인트를 호출한다. 수정 전 코드에서는 실패하도록 짰다.
실행: <venv>/bin/python -m pytest tests/test_audit_mid_b_2026_10.py -q
"""
import io
import os
import queue
import sys
import uuid
import zipfile
from datetime import datetime, timedelta

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))
os.environ.setdefault("GLHAC_DB_URL", "sqlite:///./glhac_auditmid_b_test.db")
os.environ.setdefault("GLHAC_DEV", "1")
os.environ.setdefault("GLHAC_NOTIFY_WORKER", "0")
os.environ.setdefault("GLHAC_RATE_LIMIT_PER_MIN", "0")

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from starlette.requests import Request  # noqa: E402

from app import auth, intake, models, notify  # noqa: E402
from app import main as m  # noqa: E402
from app.db import SessionLocal  # noqa: E402
from app.main import app  # noqa: E402


def _uniq():
    return uuid.uuid4().hex[:10]


def _req(xff=None, client=("9.9.9.9", 1234)):
    headers = []
    if xff is not None:
        headers.append((b"x-forwarded-for", xff.encode()))
    return Request({"type": "http", "headers": headers, "client": client,
                    "method": "GET", "path": "/", "query_string": b""})


def _login(c, u, p="pw"):
    r = c.post("/auth/login", json={"username": u, "password": p})
    assert r.status_code == 200, r.text
    return {"Authorization": "Bearer " + r.json()["token"]}


# ── MB1. 공개 게시판 노출 ────────────────────────────────────────────────
def _board_post(c, title, contact, password="mbpwabcdefgh"):
    r = c.post("/board/posts", json={"title": title, "body": "본문 비밀 내용입니다 열 글자 넘김",
                                     "author_name": "홍길동", "contact": contact,
                                     "password": password})
    assert r.status_code in (200, 201), r.text
    return r.json()


def test_mb1_public_list_has_no_title():
    with TestClient(app) as c:
        tag = "공개유출제목" + _uniq()
        _board_post(c, tag, "010-8123-%04d" % (int(_uniq(), 16) % 10000))
        r = c.get("/board/public", params={"size": 50})
        assert r.status_code == 200
        assert tag not in r.text
        assert r.json()["items"] and all("title" not in x for x in r.json()["items"])


def test_mb1_find_without_password_has_no_title_but_with_password_has():
    with TestClient(app) as c:
        tag = "찾기제목" + _uniq()
        contact = "010-7%03d-%04d" % (int(_uniq()[:4], 16) % 1000, int(_uniq()[4:8], 16) % 10000)
        _board_post(c, tag, contact, "findpwabcdef")
        r = c.post("/board/find", json={"contact": contact})
        assert r.status_code == 200
        assert tag not in r.text
        assert all("title" not in x for x in r.json())
        assert r.json()[0]["post_no"] and r.json()[0]["created_at"]
        ok = c.post("/board/find", json={"contact": contact, "password": "findpwabcdef"})
        assert ok.json()[0]["title"] == tag
        bad = c.post("/board/find", json={"contact": contact, "password": "wrongpwabcd"})
        assert bad.status_code == 401 and tag not in bad.text


# ── MB2. IP 위조 ─────────────────────────────────────────────────────────
def test_mb2_extract_ip_uses_last_xff_value():
    assert m._extract_real_ip(_req("1.1.1.1, 2.2.2.2, 3.3.3.3")) == "3.3.3.3"   # 앞쪽은 위조 가능
    assert m._extract_real_ip(_req("5.5.5.5")) == "5.5.5.5"
    assert m._extract_real_ip(_req()) == "9.9.9.9"
    assert m._extract_real_ip(_req("  ")) == "9.9.9.9"
    assert m._extract_real_ip(_req("1.1.1.1,")) == "9.9.9.9"   # 끝이 비면 소켓 주소


def test_mb2_all_ip_consumers_share_helper():
    # 같은 실제 접속 IP 면 앞쪽 위조값이 달라도 같은 식별자, 다르면 다른 식별자
    a = m._ip_hash(_req("aaa, 7.7.7.7"))
    b = m._ip_hash(_req("bbb, 7.7.7.7"))
    c_ = m._ip_hash(_req("aaa, 8.8.8.8"))
    assert a == b and a != c_
    # 쓰기 레이트리밋 액터 키도 동일 헬퍼
    assert m._rl_actor_key(_req("zzz, 7.7.7.7")) == "ip:7.7.7.7"


# ── MB3. IP 단위 로그인 실패 제한 + check-username ──────────────────────
def test_mb3_login_ip_limit_blocks_across_usernames(monkeypatch):
    monkeypatch.setenv("GLHAC_LOGIN_IP_LIMIT", "3")
    ip = "10.%d.%d.1" % (int(_uniq()[:2], 16), int(_uniq()[2:4], 16))
    with TestClient(app) as c:
        for i in range(3):   # 사용자명을 매번 바꿔 username 단위 제한은 못 걸리게
            r = c.post("/auth/login", json={"username": "nobody" + _uniq(), "password": "x"},
                       headers={"X-Forwarded-For": "spoof%d, %s" % (i, ip)})
            assert r.status_code == 401
        r = c.post("/auth/login", json={"username": "nobody" + _uniq(), "password": "x"},
                   headers={"X-Forwarded-For": "other, " + ip})
        assert r.status_code == 429
        # 올바른 계정이어도 같은 IP 는 막힌다
        r = c.post("/auth/login", json={"username": "operator1", "password": "pw"},
                   headers={"X-Forwarded-For": ip})
        assert r.status_code == 429
        # 다른 IP 는 영향 없음
        r = c.post("/auth/login", json={"username": "operator1", "password": "pw"},
                   headers={"X-Forwarded-For": "x, 10.250.250.250"})
        assert r.status_code == 200


def test_mb3_login_ip_limit_zero_disables(monkeypatch):
    monkeypatch.setenv("GLHAC_LOGIN_IP_LIMIT", "0")
    ip = "10.77.%d.2" % (int(_uniq()[:2], 16))
    with TestClient(app) as c:
        for _ in range(8):
            r = c.post("/auth/login", json={"username": "nobody" + _uniq(), "password": "x"},
                       headers={"X-Forwarded-For": ip})
            assert r.status_code == 401


def test_mb3_check_username_20_per_minute_per_ip():
    ip = "10.99.%d.3" % (int(_uniq()[:2], 16))
    with TestClient(app) as c:
        for i in range(20):
            r = c.get("/auth/check-username", params={"username": "u%s%d" % (_uniq(), i)},
                      headers={"X-Forwarded-For": ip})
            assert r.status_code == 200, i
        r = c.get("/auth/check-username", params={"username": "extra"}, headers={"X-Forwarded-For": ip})
        assert r.status_code == 429
        # 가입 화면은 막지 않는다 — 다른 IP 는 정상
        r = c.get("/auth/check-username", params={"username": "extra"},
                  headers={"X-Forwarded-For": "10.99.250.250"})
        assert r.status_code == 200


# ── MB4. SMTP STARTTLS ───────────────────────────────────────────────────
class _FakeSMTP:
    instances = []

    def __init__(self, host, port, timeout=None, tls_fail=False):
        self.calls = []
        self.tls_fail = tls_fail
        _FakeSMTP.instances.append(self)

    def starttls(self):
        self.calls.append("starttls")
        if self.tls_fail:
            raise RuntimeError("tls down")

    def login(self, u, p):
        self.calls.append("login")

    def sendmail(self, *a):
        self.calls.append("sendmail")

    def quit(self):
        self.calls.append("quit")

    def close(self):
        self.calls.append("close")


def _smtp_env(monkeypatch):
    monkeypatch.setenv("GLHAC_SMTP_HOST", "smtp.test")
    monkeypatch.setenv("GLHAC_SMTP_USER", "u")
    monkeypatch.setenv("GLHAC_SMTP_PASS", "p")
    monkeypatch.setenv("GLHAC_SMTP_FROM", "f@test")
    monkeypatch.delenv("GLHAC_SMTP_SSL", raising=False)
    _FakeSMTP.instances = []


def test_mb4_starttls_failure_does_not_login(monkeypatch):
    import smtplib
    _smtp_env(monkeypatch)

    class TlsFail(_FakeSMTP):
        def __init__(self, *a, **k):
            super().__init__(*a, tls_fail=True, **k)
    monkeypatch.setattr(smtplib, "SMTP", TlsFail)
    res = notify._send_email({"email": "to@test"}, "본문")
    assert res == {"channel": "email", "ok": False, "reason": "smtp_tls_failed"}
    calls = _FakeSMTP.instances[0].calls
    assert "login" not in calls and "sendmail" not in calls


def test_mb4_starttls_ok_logs_in_and_sends(monkeypatch):
    import smtplib
    _smtp_env(monkeypatch)
    monkeypatch.setattr(smtplib, "SMTP", _FakeSMTP)
    res = notify._send_email({"email": "to@test"}, "본문")
    assert res["ok"] is True
    assert _FakeSMTP.instances[0].calls[:3] == ["starttls", "login", "sendmail"]


def test_mb4_smtp_ssl_flag_uses_smtp_ssl(monkeypatch):
    import smtplib
    _smtp_env(monkeypatch)
    monkeypatch.setenv("GLHAC_SMTP_SSL", "1")
    used = []

    class Ssl(_FakeSMTP):
        def __init__(self, *a, **k):
            used.append("ssl")
            super().__init__(*a, **k)

    def _plain(*a, **k):
        raise AssertionError("SMTP_SSL 모드에서 평문 SMTP 를 쓰면 안 된다")
    monkeypatch.setattr(smtplib, "SMTP_SSL", Ssl)
    monkeypatch.setattr(smtplib, "SMTP", _plain)
    res = notify._send_email({"email": "to@test"}, "본문")
    assert res["ok"] is True and used == ["ssl"]
    assert "starttls" not in _FakeSMTP.instances[0].calls


# ── MB5. LLM auto 는 로컬만 ──────────────────────────────────────────────
def test_mb5_auto_never_falls_back_to_remote(monkeypatch):
    from app import ai_local
    monkeypatch.setattr(ai_local, "PROVIDER", "auto")
    monkeypatch.setattr(ai_local, "OPENAI_KEY", "sk-test")
    monkeypatch.setattr(ai_local, "_ollama_up", lambda timeout=2.5: False)
    assert ai_local.provider(refresh=True) == "none"
    monkeypatch.setattr(ai_local, "_ollama_up", lambda timeout=2.5: True)
    assert ai_local.provider(refresh=True) == "ollama"
    monkeypatch.setattr(ai_local, "PROVIDER", "openai")
    assert ai_local.provider(refresh=True) == "openai"   # 명시했을 때만 원격


# ── MB6. PSrE·TSA 응답 검증 ──────────────────────────────────────────────
class _Resp:
    def __init__(self, status=200, body=None, json_ok=True):
        self.status_code = status
        self._body = body
        self._json_ok = json_ok
        self.text = ""

    def json(self):
        if not self._json_ok:
            raise ValueError("not json")
        return self._body


def _psre_env(monkeypatch):
    monkeypatch.setenv("GLHAC_PSRE_PROVIDER", "vida")
    monkeypatch.setenv("GLHAC_PSRE_API_URL", "https://psre.test/sign")
    monkeypatch.setenv("GLHAC_PSRE_TOKEN", "tok")


@pytest.mark.parametrize("resp", [
    _Resp(200, None, json_ok=False),      # JSON 아님
    _Resp(200, ["a"]),                    # JSON 이지만 객체 아님
    _Resp(200, {}),                       # 서명값 없음
    _Resp(200, {"signer": "x"}),          # 서명값 없음
    _Resp(500, {"id": "x"}),              # HTTP 오류
])
def test_mb6_psre_rejects_2xx_without_signature_value(monkeypatch, resp):
    import httpx
    _psre_env(monkeypatch)
    monkeypatch.setattr(httpx, "post", lambda *a, **k: resp)
    assert m._psre_sign("h" * 64, "signer", "doc") is None


def test_mb6_psre_accepts_real_signature_value(monkeypatch):
    import httpx
    _psre_env(monkeypatch)
    monkeypatch.setattr(httpx, "post", lambda *a, **k: _Resp(200, {"signature_ref": "SIG-1"}))
    res = m._psre_sign("h" * 64, "signer", "doc")
    assert res and res["signature_ref"] == "SIG-1"


def test_mb6_certified_event_only_with_signature_value(monkeypatch):
    import httpx
    _psre_env(monkeypatch)
    db = SessionLocal()
    try:
        case = models.CaseApplication(company_name="PSrE" + _uniq(), org_id="org_demo", status="draft")
        db.add(case)
        db.flush()

        def n_events():
            return (db.query(models.WorkflowEvent)
                    .filter_by(case_id=case.case_id, action=m._PSRE_CERTIFIED_EVENT).count())

        monkeypatch.setattr(httpx, "post", lambda *a, **k: _Resp(200, {}))
        assert m._psre_certify(db, case, "certificate", "u", "ref", "h" * 64) is None
        assert n_events() == 0
        monkeypatch.setattr(httpx, "post", lambda *a, **k: _Resp(200, {"id": "D-9"}))
        assert m._psre_certify(db, case, "certificate", "u", "ref", "h" * 64) is not None
        assert n_events() == 1
        db.rollback()
    finally:
        db.close()


@pytest.mark.parametrize("resp", [
    _Resp(200, None, json_ok=False),
    _Resp(200, {}),
    _Resp(200, {"timestamp": "2026-10-01T00:00:00"}),   # 시각만 있고 스탬프 식별값 없음
    _Resp(200, {"serial": "S1"}),                       # 식별값만 있고 시각 없음(우리 시계로 채우면 가짜)
])
def test_mb6_tsa_rejects_incomplete_response(monkeypatch, resp):
    import httpx
    monkeypatch.setenv("GLHAC_TSA_URL", "https://tsa.test/ts")
    monkeypatch.setattr(httpx, "post", lambda *a, **k: resp)
    assert m._tsa_stamp("d" * 64) is None
    assert m._tsa_meta("d" * 64) == {}


def test_mb6_tsa_accepts_complete_response(monkeypatch):
    import httpx
    monkeypatch.setenv("GLHAC_TSA_URL", "https://tsa.test/ts")
    monkeypatch.setattr(httpx, "post",
                        lambda *a, **k: _Resp(200, {"timestamp": "2026-10-01T00:00:00", "serial": "S1"}))
    assert m._tsa_meta("d" * 64) == {"tsa_time": "2026-10-01T00:00:00", "tsa_ref": "S1", "tsa": "connected"}


# ── MB7. 공개 인증서 검증 signature_status ──────────────────────────────
def _make_cert(signed_at, tamper):
    db = SessionLocal()
    try:
        case = models.CaseApplication(company_name="Cert" + _uniq(), org_id="org_demo", status="draft")
        db.add(case)
        db.flush()
        cert = models.HalalCertificate(case_id=case.case_id, certificate_no="CN-" + _uniq(),
                                       scope=["P"], issue_date=signed_at.date().isoformat(),
                                       expiry_date="2099-01-01", status="active", qr_token="qr" + _uniq())
        db.add(cert)
        db.flush()
        ph, sigv = m._sign_payload(m._cert_canonical(cert))
        db.add(models.Signature(case_id=case.case_id, subject_type="certificate", subject_id=cert.id,
                                signer="t", payload_hash=ph,
                                signature_value=("bad" + sigv) if tamper else sigv, signed_at=signed_at))
        db.commit()
        return cert.qr_token
    finally:
        db.close()


def test_mb7_signature_status_three_way(monkeypatch):
    monkeypatch.setenv("GLHAC_AUDIT_CHAIN_EPOCH", "2026-09-10T12:00:00")
    before = datetime(2026, 9, 1, 9, 0)
    after = datetime(2026, 9, 20, 9, 0)
    with TestClient(app) as c:
        r = c.get("/verify/" + _make_cert(after, tamper=False)).json()
        assert r["signature_status"] == "valid" and r["signature_valid"] is True
        r = c.get("/verify/" + _make_cert(before, tamper=True)).json()
        assert r["signature_status"] == "legacy" and r["signature_valid"] is False
        r = c.get("/verify/" + _make_cert(after, tamper=True)).json()
        assert r["signature_status"] == "invalid" and r["signature_valid"] is False
        # 경계 직전/직후(같은 날 오전 vs 오후)
        assert c.get("/verify/" + _make_cert(datetime(2026, 9, 10, 11, 59), True)).json()["signature_status"] == "legacy"
        assert c.get("/verify/" + _make_cert(datetime(2026, 9, 10, 12, 1), True)).json()["signature_status"] == "invalid"


def test_mb7_epoch_env_moves_boundary(monkeypatch):
    monkeypatch.setenv("GLHAC_AUDIT_CHAIN_EPOCH", "2027-01-01T00:00:00")
    with TestClient(app) as c:
        tok = _make_cert(datetime(2026, 12, 31, 9, 0), tamper=True)
        assert c.get("/verify/" + tok).json()["signature_status"] == "legacy"


# ── MB8. SiHALAL 신원 '검증' 은 컨설턴트 확인 ────────────────────────────
def test_mb8_identity_verify_marks_consultant_match():
    with TestClient(app) as c:
        h = _login(c, "consultant1")
        case = models.CaseApplication  # noqa: F841
        db = SessionLocal()
        try:
            cs = models.CaseApplication(company_name="SH" + _uniq(), org_id="org_demo", status="draft")
            db.add(cs)
            db.commit()
            cid = cs.case_id
        finally:
            db.close()
        ei = c.post(f"/cases/{cid}/sihalal/identity/link", json={"external_email": "mb8@x.com"},
                    headers=h).json()
        eid = ei["external_identity_id"]
        ok = c.post(f"/sihalal/identity/{eid}/verify", json={"expected_identifier": "mb8@x.com"}, headers=h)
        assert ok.status_code == 200
        body = ok.json()
        # 저장 상태값은 그대로(게이트 동작 불변)
        assert body["verification_status"] == "verified" and body["identifier_match"] is True
        assert body["method"] == "consultant_match" and body["external_verified"] is False
        bad = c.post(f"/sihalal/identity/{eid}/verify", json={"expected_identifier": "no@x.com"}, headers=h).json()
        assert bad["verification_status"] == "mismatch_found"
        assert bad["method"] == "consultant_match" and bad["external_verified"] is False
        db = SessionLocal()
        try:
            evs = (db.query(models.WorkflowEvent)
                   .filter_by(case_id=cid, action="sihalal.identity_consultant_check").all())
            assert len(evs) == 2
            for e in evs:
                assert e.payload["method"] == "consultant_match" and e.payload["external_verified"] is False
        finally:
            db.close()


# ── MB9. 업로드 상한 20MB 공유 ───────────────────────────────────────────
def test_mb9_upload_limit_is_20mb_and_shared():
    assert m._UPLOAD_MAX_BYTES == 20 * 1024 * 1024
    assert intake.UPLOAD_MAX_BYTES == m._UPLOAD_MAX_BYTES


def test_mb9_original_retention_uses_shared_constant(monkeypatch):
    monkeypatch.setattr(intake, "parse_file", lambda name, raw, ocr_sink=None: "")
    monkeypatch.setattr(intake, "classify", lambda name, text: {"doc_type": "other", "fields": {}})
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("a.txt", "x" * 50)
    res = intake.intake_zip(buf.getvalue())
    assert res["classified"][0]["content_b64"] is not None
    monkeypatch.setattr(intake, "UPLOAD_MAX_BYTES", 10)   # 상수를 바꾸면 보관 여부가 따라 바뀐다
    res = intake.intake_zip(buf.getvalue())
    assert res["classified"][0]["content_b64"] is None


def test_mb9_zip_endpoints_cap_base64_at_20mb_equivalent():
    # 유효 base64 길이(패딩 포함)로 정확히 20MB 와 20MB+1 의 경계
    ok_len = ((m._UPLOAD_MAX_BYTES + 2) // 3) * 4
    too_big = "A" * (ok_len + 8)
    with TestClient(app) as c:
        h = _login(c, "applicant1")
        cid = c.post("/cases", json={"company_name": "Z" + _uniq()}, headers=h).json()["case_id"]
        for ep in ("intake-zip", "intake-zip-stream"):
            r = c.post(f"/cases/{cid}/{ep}", json={"zip_b64": too_big}, headers=h)
            assert r.status_code == 413, (ep, r.status_code)
            assert "20MB" in r.text
            # 상한 이내는 크기 검사를 통과(압축 파일이 아니라서 다른 오류)
            r = c.post(f"/cases/{cid}/{ep}", json={"zip_b64": "A" * 100}, headers=h)
            assert r.status_code != 413


# ── MB10. intake.py 한계·표·ZIP·로그 ─────────────────────────────────────
def _pdf(pages):
    import fitz
    doc = fitz.open()
    for i in range(pages):
        pg = doc.new_page()
        pg.insert_text((72, 72), "PAGEMARK%03d lorem ipsum dolor sit amet" % i)
    return doc.tobytes()


def test_mb10_pdf_reads_40_pages_not_8():
    txt = intake.parse_file("big.pdf", _pdf(45))
    assert "PAGEMARK039" in txt          # 40번째 쪽까지
    assert "PAGEMARK040" not in txt      # 41번째부터 제외
    assert "PAGEMARK010" in txt          # 예전 8쪽 한계 밖


def test_mb10_pdf_text_cap(monkeypatch):
    monkeypatch.setattr(intake, "PDF_MAX_TEXT_CHARS", 300)
    txt = intake.parse_file("big.pdf", _pdf(30))
    assert len(txt) <= 300
    assert "PAGEMARK029" not in txt


def test_mb10_docx_table_cells_extracted_by_row():
    import docx
    d = docx.Document()
    d.add_paragraph("본문단락")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "원료A"
    t.cell(0, 1).text = "공급사X"
    t.cell(1, 0).text = "원료B"
    t.cell(1, 1).text = "공급사Y"
    buf = io.BytesIO()
    d.save(buf)
    txt = intake.parse_file("t.docx", buf.getvalue())
    lines = txt.splitlines()
    assert "본문단락" in lines
    assert any("원료A" in ln and "공급사X" in ln for ln in lines)   # 행 단위
    assert any("원료B" in ln and "공급사Y" in ln for ln in lines)
    assert not any("원료A" in ln and "원료B" in ln for ln in lines)


def test_mb10_zip_over_200_entries_reports_skipped(monkeypatch, caplog):
    monkeypatch.setattr(intake, "parse_file", lambda name, raw, ocr_sink=None: "")
    monkeypatch.setattr(intake, "classify", lambda name, text: {"doc_type": "other", "fields": {}})
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for i in range(205):
            z.writestr("f%03d.txt" % i, "x")
    files = intake.extract_zip(buf.getvalue())
    assert len(files) == 200 and files.skipped == 5
    with caplog.at_level("WARNING"):
        res = intake.intake_zip(buf.getvalue())
    assert res["skipped_entries"] == 5 and len(res["classified"]) == 200
    assert any("초과" in r.getMessage() for r in caplog.records)
    # 200개 이하면 0
    buf2 = io.BytesIO()
    with zipfile.ZipFile(buf2, "w") as z:
        z.writestr("only.txt", "x")
    assert intake.intake_zip(buf2.getvalue())["skipped_entries"] == 0


def test_mb10_corrupt_file_logs_warning(caplog):
    with caplog.at_level("WARNING"):
        assert intake.parse_file("broken.pdf", b"%PDF-1.4 not really a pdf") == ""
        assert intake.parse_file("broken.docx", b"not a zip") == ""
    msgs = [r.getMessage() for r in caplog.records]
    assert any("broken.pdf" in s for s in msgs) and any("broken.docx" in s for s in msgs)


# ── MB11. 자동 파싱 큐 ───────────────────────────────────────────────────
def _doc(db, confidence=0.9):
    case = models.CaseApplication(company_name="Q" + _uniq(), org_id="org_demo", status="draft")
    db.add(case)
    db.flush()
    d = models.DocumentAsset(case_id=case.case_id, filename="a.pdf", doc_type="other",
                             confidence=confidence, content_b64="JVBERi0xLjQK")
    db.add(d)
    db.commit()
    return d.document_id


def test_mb11_full_queue_leaves_doc_pending_and_dedups(monkeypatch):
    q = queue.Queue(maxsize=1)
    monkeypatch.setattr(m, "_AUTOPARSE", True)
    monkeypatch.setattr(m, "_autoparse_q", q)
    monkeypatch.setattr(m, "_autoparse_inflight", set())
    db = SessionLocal()
    try:
        a, b = _doc(db), _doc(db, confidence=0.9)
        assert m.autoparse_enqueue(a) is True
        assert m.autoparse_enqueue(a) is False        # 같은 문서 중복 투입 방지
        assert q.qsize() == 1
        # 큐가 가득 참: 막히지 않고 즉시 False, 문서는 대기 상태(confidence NULL)
        assert m.autoparse_enqueue(b) is False
        db.expire_all()
        assert db.get(models.DocumentAsset, b).confidence is None
        assert b not in m._autoparse_inflight         # 다시 시도할 수 있다
        assert db.get(models.DocumentAsset, a).confidence == 0.9   # 큐에 들어간 문서는 건드리지 않음
    finally:
        db.close()


class _OneShotQ(queue.Queue):
    """첫 get 은 항목을 주고 두 번째 get 에서 워커 루프를 끝낸다."""
    def __init__(self, *a, **k):
        super().__init__(*a, **k)
        self.gets = 0

    def get(self, *a, **k):
        self.gets += 1
        if self.gets > 1:
            raise KeyboardInterrupt
        return super().get(*a, **k)


@pytest.mark.parametrize("fails", [False, True])
def test_mb11_worker_removes_inflight_on_done_or_failure(monkeypatch, fails):
    q = _OneShotQ(maxsize=5)
    monkeypatch.setattr(m, "_AUTOPARSE", True)
    monkeypatch.setattr(m, "_autoparse_q", q)
    monkeypatch.setattr(m, "_autoparse_inflight", set())

    def fake(db, d, c, keep_doc_type=False, **k):
        if fails:
            raise RuntimeError("boom")
    monkeypatch.setattr(m, "_reprocess_doc", fake)
    db = SessionLocal()
    try:
        a = _doc(db)
        assert m.autoparse_enqueue(a) is True
        assert a in m._autoparse_inflight
        with pytest.raises(KeyboardInterrupt):
            m._autoparse_worker()
        assert a not in m._autoparse_inflight
        assert m.autoparse_enqueue(a) is True         # 끝났으니 다시 넣을 수 있다
    finally:
        db.close()


# ── MB12. 알림 수신자 오배정 ─────────────────────────────────────────────
def _capture_providers(monkeypatch):
    seen = []

    def make(ch):
        def fn(contacts, text, notification=None):
            seen.append((ch, dict(contacts)))
            if not contacts.get("phone" if ch == "sms" else "email"):
                return {"channel": ch, "ok": False, "reason": "no_contact"}
            return {"channel": ch, "ok": True}
        return fn
    monkeypatch.setitem(notify.PROVIDERS, "sms", make("sms"))
    monkeypatch.setitem(notify.PROVIDERS, "email", make("email"))
    return seen


def _notif_setup(role, with_profile=True):
    db = SessionLocal()
    try:
        org_id = "org_" + _uniq()
        cons_id = "cons_" + _uniq()
        if with_profile is not None:   # FK: 담당 컨설턴트는 실제 사용자여야 한다(프로필만 없을 수 있다)
            db.add(models.User(user_id=cons_id, username=cons_id, password_hash="x",
                               role="consultant", org_id="org_demo"))
            db.flush()
        db.add(models.Org(org_id=org_id, name="O", consultant_id=cons_id if with_profile is not None else None))
        if with_profile:
            db.add(models.ConsultantProfile(consultant_id=cons_id, phone="010-CONS-0000",
                                            email="cons@x.test", status="active"))
        case = models.CaseApplication(company_name="N" + _uniq(), org_id=org_id, status="draft",
                                      phone="010-COMP-1111", email="company@x.test", notify_consent=True)
        db.add(case)
        db.flush()
        n = models.Notification(case_id=case.case_id, org_id=org_id, role=role, event_type="t",
                                channels=["sms", "email"], title="T", body="B", status="unsent")
        db.add(n)
        db.commit()
        return n.notification_id
    finally:
        db.close()


def _drain_one(nid):
    db = SessionLocal()
    try:
        # 다른 테스트가 남긴 미발송 알림이 batch 를 차지하지 않도록 이 알림만 남긴다
        db.query(models.Notification).filter(models.Notification.notification_id != nid,
                                             models.Notification.status == "unsent"
                                             ).update({"status": "skipped"}, synchronize_session=False)
        db.commit()
        m.drain_notifications(db)
        db.expire_all()
        return db.get(models.Notification, nid)
    finally:
        db.close()


def test_mb12_consultant_notice_goes_to_consultant_not_company(monkeypatch):
    seen = _capture_providers(monkeypatch)
    n = _drain_one(_notif_setup("consultant"))
    assert n.status == "sent"
    by_ch = dict(seen)
    assert by_ch["sms"]["phone"] == "010-CONS-0000" and by_ch["email"]["email"] == "cons@x.test"
    assert all("COMP" not in str(c.get("phone")) and "company@" not in str(c.get("email")) for _, c in seen)


def test_mb12_consultant_without_profile_is_no_contact_not_company(monkeypatch):
    seen = _capture_providers(monkeypatch)
    n = _drain_one(_notif_setup("consultant", with_profile=False))
    assert n.status == "skipped"
    assert "no_contact" in (n.last_error or "")
    assert all(not c.get("phone") and not c.get("email") for _, c in seen)


def test_mb12_auditor_notice_never_goes_to_company(monkeypatch):
    seen = _capture_providers(monkeypatch)
    n = _drain_one(_notif_setup("auditor"))
    assert n.status == "skipped" and "no_contact" in (n.last_error or "")
    assert all(not c.get("phone") and not c.get("email") for _, c in seen)


def test_mb12_company_roles_still_reach_company(monkeypatch):
    for role in (None, "applicant", "client"):
        seen = _capture_providers(monkeypatch)
        n = _drain_one(_notif_setup(role))
        assert n.status == "sent", role
        assert dict(seen)["sms"]["phone"] == "010-COMP-1111"
