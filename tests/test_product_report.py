"""현장심사 보고서 제품표 — 무엇을 심사했는지 사진과 설명으로 특정한다.

배경(LPH 피드백): Site-audit report 템플릿의 제품표는 No / Name(제품명+설명) / Image
세 칸이다. 우리 템플릿에는 'Product name and description' 제목만 있고 표가 없었고,
제품에 설명 필드조차 없었다. 이름만 있는 보고서로는 심사 대상이 특정되지 않는다.

실행: <venv>/bin/python -m pytest tests/test_product_report.py -q
"""
import base64
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))
os.environ.setdefault("GLHAC_DB_URL", "sqlite:///./glhac_prodrep_test.db")

from fastapi.testclient import TestClient   # noqa: E402

import app.main as m                        # noqa: E402
from app.main import app                    # noqa: E402


def _png(color=(40, 90, 60), size=(300, 200)):
    from PIL import Image
    b = io.BytesIO()
    Image.new("RGB", size, color).save(b, "PNG")
    return base64.b64encode(b.getvalue()).decode()


def _tok(c, u="applicant1", p="pw"):
    r = c.post("/auth/login", json={"username": u, "password": p})
    assert r.status_code == 200, r.text
    return {"Authorization": "Bearer " + r.json()["token"]}


def _case(c, h):
    return c.post("/cases", json={"company_name": "PT Produk Uji"}, headers=h).json()["case_id"]


def _prod_table(doc):
    return next((t for t in doc.tables
                 if len(t.columns) == 3 and t.rows
                 and t.rows[0].cells[0].text.strip().lower() == "no"
                 and t.rows[0].cells[2].text.strip().lower().startswith("image")), None)


def _has_image(row):
    from docx.oxml.ns import qn
    return any(c._tc.findall(".//" + qn("w:drawing")) for c in row.cells)


# ── 템플릿 ──────────────────────────────────────────────────────────────
def test_템플릿에_제품표가_있다():
    """제목만 있고 표가 없어 제품이 보고서에서 통째로 빠져 있었다."""
    import docx
    doc = docx.Document(m.FACTORY_AUDIT_TEMPLATE)
    assert any(p.text.strip().startswith("Product name and description")
               for p in doc.paragraphs), "제품 제목 없음"
    assert _prod_table(doc) is not None, "제품표 없음"


# ── 설명 필드 ───────────────────────────────────────────────────────────
def test_제품에_설명을_적고_읽을_수_있다():
    with TestClient(app) as c:
        h = _tok(c)
        cid = _case(c, h)
        pid = c.post(f"/cases/{cid}/products",
                     json={"name": "Bumbu A", "description": "즉석 조미료 200g"},
                     headers=h).json()["product_id"]
        row = c.get(f"/cases/{cid}/products", headers=h).json()[0]
        assert row["description"] == "즉석 조미료 200g"
        c.patch(f"/cases/{cid}/products/{pid}", json={"description": "고침 250g"}, headers=h)
        assert c.get(f"/cases/{cid}/products/{pid}",
                     headers=h).json()["description"] == "고침 250g"


def test_보고서에_필요한_것이_빠졌는지_알려준다():
    """설명·사진이 없으면 보고서가 빈칸으로 나간다 — 내기 전에 화면이 알려줘야 한다."""
    with TestClient(app) as c:
        h = _tok(c)
        cid = _case(c, h)
        pid = c.post(f"/cases/{cid}/products", json={"name": "Bumbu B"},
                     headers=h).json()["product_id"]
        row = c.get(f"/cases/{cid}/products", headers=h).json()[0]
        assert row["report_ready"] is False
        assert set(row["missing_for_report"]) == {"description", "photo"}
        c.patch(f"/cases/{cid}/products/{pid}", json={"description": "설명"}, headers=h)
        assert c.get(f"/cases/{cid}/products",
                     headers=h).json()[0]["missing_for_report"] == ["photo"]
        c.post(f"/cases/{cid}/products/{pid}/photo",
               json={"file_b64": _png(), "filename": "p.png"}, headers=h)
        row = c.get(f"/cases/{cid}/products", headers=h).json()[0]
        assert row["report_ready"] is True and row["missing_for_report"] == []


# ── 보고서 생성 ─────────────────────────────────────────────────────────
def test_제품명과_설명이_보고서에_같이_인쇄된다():
    import docx
    with TestClient(app) as c:
        h = _tok(c)
        cid = _case(c, h)
        c.post(f"/cases/{cid}/products",
               json={"name": "SEMUR BETAWI", "description": "Bumbu instan 200g"}, headers=h)
        c0 = m.SessionLocal().get(m.models.CaseApplication, cid)
        doc = docx.Document(io.BytesIO(m._factory_audit_docx_bytes(m.SessionLocal(), c0)))
        t = _prod_table(doc)
        txt = t.rows[1].cells[1].text
        assert "SEMUR BETAWI" in txt and "Bumbu instan 200g" in txt


def test_사진이_보고서_제품표에_들어간다():
    import docx
    with TestClient(app) as c:
        h = _tok(c)
        cid = _case(c, h)
        pid = c.post(f"/cases/{cid}/products", json={"name": "Bumbu C"},
                     headers=h).json()["product_id"]
        c.post(f"/cases/{cid}/products/{pid}/photo",
               json={"file_b64": _png(), "filename": "c.png"}, headers=h)
        c0 = m.SessionLocal().get(m.models.CaseApplication, cid)
        doc = docx.Document(io.BytesIO(m._factory_audit_docx_bytes(m.SessionLocal(), c0)))
        assert _has_image(_prod_table(doc).rows[1])


def test_사진이_없으면_빈칸이_아니라_미제출로_적는다():
    """빈칸으로 두면 누락인지 원래 없는 것인지 구분되지 않는다."""
    import docx
    with TestClient(app) as c:
        h = _tok(c)
        cid = _case(c, h)
        c.post(f"/cases/{cid}/products", json={"name": "Bumbu D"}, headers=h)
        c0 = m.SessionLocal().get(m.models.CaseApplication, cid)
        doc = docx.Document(io.BytesIO(m._factory_audit_docx_bytes(m.SessionLocal(), c0)))
        cell = _prod_table(doc).rows[1].cells[2].text
        assert "미제출" in cell or "No photo" in cell


def test_첫_사진이_깨져_있으면_다음_사진을_쓴다():
    """전송 중 잘린 파일 한 장 때문에 '낸 것이 없는' 보고서가 나가면 안 된다."""
    import docx
    with TestClient(app) as c:
        h = _tok(c)
        cid = _case(c, h)
        pid = c.post(f"/cases/{cid}/products", json={"name": "Bumbu E"},
                     headers=h).json()["product_id"]
        broken = base64.b64encode(b"\x89PNG\r\n\x1a\n" + b"\x00" * 20).decode()
        c.post(f"/cases/{cid}/products/{pid}/photo",
               json={"file_b64": broken, "filename": "broken.png"}, headers=h)
        c.post(f"/cases/{cid}/products/{pid}/photo",
               json={"file_b64": _png(), "filename": "good.png"}, headers=h)
        c0 = m.SessionLocal().get(m.models.CaseApplication, cid)
        doc = docx.Document(io.BytesIO(m._factory_audit_docx_bytes(m.SessionLocal(), c0)))
        assert _has_image(_prod_table(doc).rows[1])


def test_제품_수만큼_행이_맞춰진다():
    import docx
    with TestClient(app) as c:
        h = _tok(c)
        cid = _case(c, h)
        for i in range(4):
            c.post(f"/cases/{cid}/products", json={"name": "P%d" % i}, headers=h)
        c0 = m.SessionLocal().get(m.models.CaseApplication, cid)
        doc = docx.Document(io.BytesIO(m._factory_audit_docx_bytes(m.SessionLocal(), c0)))
        t = _prod_table(doc)
        assert len(t.rows) == 5                      # 머리 1 + 제품 4
        assert [r.cells[0].text.strip() for r in t.rows[1:]] == ["1", "2", "3", "4"]


# ── 제품 삭제와 인증서 동결 범위 ────────────────────────────────────────
def _issue_frozen_cert(cid, pids):
    """이 케이스에 제품 범위를 동결한 인증서가 있는 상태를 만든다."""
    import app.models as models
    db = m.SessionLocal()
    try:
        db.add(models.HalalCertificate(case_id=cid, certificate_no="HC-TEST-1",
                                       status="active", frozen_product_ids=list(pids)))
        db.commit()
    finally:
        db.close()


def test_인증서에_동결된_제품은_그냥_지워지지_않는다():
    """인증서는 '이 제품들을 인증한다'는 문서다. 제품만 사라지면 없는 것을 가리킨다."""
    with TestClient(app) as c:
        h = _tok(c)
        cid = _case(c, h)
        pid = c.post(f"/cases/{cid}/products", json={"name": "Bumbu F"},
                     headers=h).json()["product_id"]
        _issue_frozen_cert(cid, [pid])
        r = c.delete(f"/cases/{cid}/products/{pid}", headers=h)
        assert r.status_code == 409, r.status_code
        d = r.json()["detail"]
        assert d["code"] == "CERT_SCOPE_FROZEN"
        assert "HC-TEST-1" in d["certificates"]        # 어느 인증서인지 알려준다
        assert c.get(f"/cases/{cid}/products", headers=h).json(), "제품이 지워졌다"


def test_동결된_제품_강제삭제는_운영자만():
    with TestClient(app) as c:
        h = _tok(c)
        cid = _case(c, h)
        pid = c.post(f"/cases/{cid}/products", json={"name": "Bumbu G"},
                     headers=h).json()["product_id"]
        _issue_frozen_cert(cid, [pid])
        assert c.delete(f"/cases/{cid}/products/{pid}?force=true",
                        headers=h).status_code == 403


def test_강제삭제하면_인증서_동결목록에서도_빠진다():
    import app.models as models
    with TestClient(app) as c:
        h = _tok(c)
        cid = _case(c, h)
        keep = c.post(f"/cases/{cid}/products", json={"name": "Keep"},
                      headers=h).json()["product_id"]
        drop = c.post(f"/cases/{cid}/products", json={"name": "Drop"},
                      headers=h).json()["product_id"]
        _issue_frozen_cert(cid, [keep, drop])
        op = _tok(c, "operator1", "pw")
        r = c.delete(f"/cases/{cid}/products/{drop}?force=true", headers=op)
        assert r.status_code == 200, r.text
        assert r.json()["removed_from_certificates"] == ["HC-TEST-1"]
        db = m.SessionLocal()
        try:
            ct = db.query(models.HalalCertificate).filter_by(case_id=cid).first()
            assert ct.frozen_product_ids == [keep]     # 남은 제품만
        finally:
            db.close()


def test_제품_삭제가_감사기록에_남는다():
    """무엇을 지웠는지 남지 않으면 되짚을 수 없다 — 인증 범위가 조용히 줄어든다."""
    import app.models as models
    with TestClient(app) as c:
        h = _tok(c)
        cid = _case(c, h)
        pid = c.post(f"/cases/{cid}/products",
                     json={"name": "Daftar Bahan Halal"}, headers=h).json()["product_id"]
        c.delete(f"/cases/{cid}/products/{pid}", headers=h)
        db = m.SessionLocal()
        try:
            logs = [a for a in db.query(models.AuditLog).all()
                    if a.action == "product.delete" and a.resource_id == pid]
            assert logs, "감사기록 없음"
            assert (logs[-1].meta or {}).get("name") == "Daftar Bahan Halal"
        finally:
            db.close()


def test_제품_사진을_지울_수_있다():
    """잘못 올린 사진이 그대로 보고서에 인쇄된다 — 되돌릴 통로가 있어야 한다."""
    with TestClient(app) as c:
        h = _tok(c)
        cid = _case(c, h)
        pid = c.post(f"/cases/{cid}/products", json={"name": "Bumbu H"},
                     headers=h).json()["product_id"]
        c.post(f"/cases/{cid}/products/{pid}/photo",
               json={"file_b64": _png(), "filename": "wrong.png"}, headers=h)
        did = c.get(f"/cases/{cid}/products/{pid}/photos", headers=h).json()[0]["document_id"]
        assert c.delete(f"/cases/{cid}/products/{pid}/photos/{did}",
                        headers=h).status_code == 200
        assert c.get(f"/cases/{cid}/products", headers=h).json()[0]["photo_count"] == 0


# ── 언어 ────────────────────────────────────────────────────────────────
def _tables(cid, lang):
    import docx
    c0 = m.SessionLocal().get(m.models.CaseApplication, cid)
    d = docx.Document(io.BytesIO(m._factory_audit_docx_bytes(m.SessionLocal(), c0, lang)))
    return d, _prod_table(d), next(t for t in d.tables if len(t.columns) == 7)


def test_보고서_값이_언어를_따른다():
    """LPH 심사원이 읽는 것은 값이다 — 사진 미제출·원재료 유형이 언어별로 나와야 한다."""
    with TestClient(app) as c:
        h = _tok(c)
        cid = _case(c, h)
        c.post(f"/cases/{cid}/products", json={"name": "Bumbu I"}, headers=h)
        c.post(f"/cases/{cid}/materials",
               json={"name": "Gula", "mat_type": "BAHAN BAKU"}, headers=h)
        want = {"ko": ("사진 미제출", "원료"),
                "id": ("Foto belum diserahkan", "Bahan Baku"),
                "en": ("No photo submitted", "Raw Material")}
        for lang, (photo, mtype) in want.items():
            _, pt, mt = _tables(cid, lang)
            assert photo in pt.rows[1].cells[2].text, (lang, pt.rows[1].cells[2].text)
            assert mt.rows[1].cells[2].text.strip() == mtype, (lang, mt.rows[1].cells[2].text)


def test_서식_라벨은_언어와_무관하게_그대로다():
    """표 머리·항목명은 템플릿 소유다. 번역하면 LPH가 받는 서식이 달라진다."""
    with TestClient(app) as c:
        h = _tok(c)
        cid = _case(c, h)
        c.post(f"/cases/{cid}/products", json={"name": "Bumbu J"}, headers=h)
        for lang in ("ko", "id", "en"):
            _, pt, mt = _tables(cid, lang)
            assert [x.text.strip() for x in pt.rows[0].cells] == ["No", "Name", "Image"]
            assert mt.rows[0].cells[1].text.strip().startswith("Name & Brand")
