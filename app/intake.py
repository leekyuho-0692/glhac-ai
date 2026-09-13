"""Document Intake AI — 설계 Part 5: ZIP → 파싱(OCR/PDF/DOCX) → gemma3 분류 → 필수서류 라우팅."""
import base64
import io
import os
import re
import zipfile
from . import ai_local

_MIME = {"pdf": "application/pdf", "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
         "bmp": "image/bmp", "tiff": "image/tiff", "tif": "image/tiff", "webp": "image/webp",
         "txt": "text/plain; charset=utf-8", "csv": "text/csv",
         "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
         "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
         "mp4": "video/mp4", "mov": "video/quicktime", "webm": "video/webm",
         "m4v": "video/x-m4v"}


def _ctype(name):
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    return _MIME.get(ext, "application/octet-stream")

# 할랄인증 서류 종류 (BPJPH/SIHALAL 기준)
DOC_TYPES = [
    "nib_business_license",     # NIB/사업자등록증
    "factory_registration",     # 공장등록증
    "product_list",             # 제품 목록
    "process_flow",             # 공정 흐름도
    "halal_certificate",        # 공급사 할랄 인증서
    "material_list",            # 원재료 목록
    "coa_msds",                 # CoA / MSDS / 성분 명세
    "supplier_declaration",     # 공급사 선언서
    "quality_cert",             # 품질/식품안전 인증(HACCP·ISO·GMP·FSSC) — 할랄 인증 아님
    "sjph_manual",              # SJPH/HPAS 매뉴얼
    # ── 물류(jasa logistik) 고유 서류 ──
    "logistics_scope",          # 취급 화물·서비스 범위(어떤 제품군을, 어떤 jasa로)
    "warehouse_layout",         # 창고 배치도 + Halal Zone 구획
    "vehicle_list",             # 차량·컨테이너 목록(인증 대상은 서비스지만 관리대상)
    "cleaning_sop",             # 세척 SOP(차량·컨테이너·창고) + Sertu 절차
    "other",
]
# 서류명 3개 언어 — 도메인 사전(domain_dict.json)의 DOC 축에서 파생한다.
# 전에는 여기 세 벌을 손으로 유지했는데, 같은 말을 사전과 코드가 따로 갖고 있어
# 인니어 원문 표기('Diagram alir proses produksi')를 서류 유형으로 잇지 못했다.
from .domain_dict import doc_labels as _doc_labels   # noqa: E402

DOC_KO = _doc_labels("ko")
DOC_EN = _doc_labels("en")
DOC_ID = _doc_labels("id")
DOC_NAME_L10N = {"en": DOC_EN, "id": DOC_ID}
# 신청 필수 서류 — 경로별로 다르다.
#
# 왜 나누는가: 인도네시아 소규모 자기선언(SEHATI)은 NIB 하나로 사업자·시설을 갈음하고,
# 공급사 할랄 인증서는 사본 대신 인증번호를 BPJPH가 직접 대조한다. 한국 정규 경로 기준
# 목록을 그대로 적용하면 인니 완성본이 '서류 부족'으로 잘못 표시된다(CV CITRA 실측).
# 목록에서 뺀다고 검토를 건너뛰는 것이 아니라, 그 경로에서 '요구되지 않는 서류'를
# 미제출로 세지 않는 것이다.
REQUIRED_DOCS = ["nib_business_license", "factory_registration", "product_list",
                 "process_flow", "material_list", "halal_certificate", "sjph_manual"]

# 물류(jasa logistik) 필수 서류 — 제품과 다른 세트다.
# 제품의 원재료·배합·공정 축이 없고, 대신 시설·차량·세척·취급범위가 핵심이다.
# BPJPH는 차량이 아니라 '물류 서비스'를 인증하므로, 차량 목록은 관리대상 증빙이다(2024-09-05).
REQUIRED_DOCS_LOGISTICS = ["nib_business_license", "sjph_manual",
                           "logistics_scope", "warehouse_layout", "vehicle_list", "cleaning_sop"]

# 경로별 필수 목록. 미정(undetermined)은 정규 기준을 쓴다 — 넓게 요구하는 쪽이 안전하다.
REQUIRED_DOCS_BY_PATHWAY = {
    "reguler": REQUIRED_DOCS,
    "self_declare": ["nib_business_license", "product_list", "process_flow",
                     "material_list", "halal_certificate", "sjph_manual"],
}
# 서류 원본 대신 다른 근거로 충족할 수 있는 항목 — 무엇으로 갈음했는지 반드시 표기한다.
DOC_ALT_SATISFY = {
    "self_declare": {
        "halal_certificate": {"by": "supplier_cert_no",
                              "note": "인도네시아 자기선언 경로 — 공급사 인증번호로 갈음(BPJPH 대조)"},
    },
}
# 그 경로에서 요구되지 않는 서류 — 화면에 '해당 없음'으로 표시해 누락과 구분한다.
DOC_NOT_APPLICABLE = {
    "self_declare": {"factory_registration":
                     "소규모 자기선언 — NIB로 갈음(별도 공장등록증 요구 없음)"},
}


# 관할·규모에서 오는 면제 — 경로와는 다른 축이다.
# 실측(CV. CITRA PRATAMA): 육류·가금 원료 때문에 판정이 reguler로 나와도, 인도네시아
# 소규모 사업자에게 '공장등록증'이라는 서류 자체가 존재하지 않는다. NIB(사업자번호)가
# 사업자 등록과 시설 등록을 겸한다. 그래서 경로 축만으로는 이 면제를 표현할 수 없다.
DOC_JURISDICTION = {
    ("ID", True): {"factory_registration":
                   "인도네시아 소규모 사업자 — NIB가 시설 등록을 겸함(별도 공장등록증 제도 없음)"},
}
_COUNTRY_ID = {"id", "idn", "indonesia", "republic of indonesia", "인도네시아"}


def _country_key(country):
    """국가 표기 흔들림 흡수. 모르는 나라는 None — 면제를 적용하지 않는다(넓은 쪽)."""
    c = (country or "").strip().lower()
    return "ID" if c in _COUNTRY_ID else None


def required_docs(pathway=None):
    """경로별 필수 서류 목록. 모르는 경로는 정규 기준(넓은 쪽)."""
    return REQUIRED_DOCS_BY_PATHWAY.get((pathway or "").lower(), REQUIRED_DOCS)


def doc_requirements(pathway=None, country=None, is_msme=None, scheme="product"):
    """이 신청 건에 실제로 요구되는 서류 — 인증 종류(scheme)·경로(pathway)·관할·규모를 함께 본다.

    scheme 이 가장 바깥 축이다. 물류(logistics)는 제품과 서류 세트 자체가 다르므로
    pathway 분기 이전에 갈라진다. 제품(product)은 기존 pathway 로직을 그대로 탄다.

    돌려주는 것
      required        요구되는 doc_type 목록
      not_applicable  {doc_type: 사유} — 요구되지 않는 서류. 목록에서 지우지 않고
                      사유와 함께 남긴다. 조용히 사라지면 면제인지 누락인지 알 수 없다.
      alt             {doc_type: {by, note}} — 원본 서류 대신 다른 근거로 충족 가능한 항목
    """
    if (scheme or "product").lower() == "logistics":
        # 물류는 경로(self_declare/reguler)에 따른 서류 차이가 아직 규범으로 확정되지
        # 않았다(BPJPH jasa logistik SJPH 매뉴얼 미공개) → 단일 세트로 시작한다.
        return {"required": list(REQUIRED_DOCS_LOGISTICS), "not_applicable": {}, "alt": {}}
    pw = (pathway or "").lower()
    req = list(required_docs(pw))
    na = dict(DOC_NOT_APPLICABLE.get(pw, {}))
    na.update(DOC_JURISDICTION.get((_country_key(country), bool(is_msme)), {}))
    req = [d for d in req if d not in na]
    alt = {k: v for k, v in DOC_ALT_SATISFY.get(pw, {}).items() if k in req}
    return {"required": req, "not_applicable": na, "alt": alt}

# 필수 서류별 요구 내용(보완 안내용) — 무엇이 담겨야 하는지 상세 설명
DOC_REQUIREMENT = {
    "nib_business_license": "사업자등록증(NIB) 스캔본 — 회사명·NIB 번호·사업장 주소가 판독 가능해야 함",
    "factory_registration": "공장등록증 — 공장 등록번호와 공장 소재지 주소 포함",
    "product_list": "인증 대상 전(全) 제품 목록 — 제품명·분류·등록유형(신규/기존 등)",
    "process_flow": "제조 공정 흐름도 — 원료입고→배합→가열→충전→포장 등 단계 순서",
    "material_list": "전(全) 원재료 목록 — 원재료명·공급사·할랄 상태(인증/선언)",
    "halal_certificate": "임계 원재료 공급사가 받은 할랄 인증서 사본 — 이 플랫폼이 발급하는 인증서가 아니라 신청자가 제출하는 입력 서류다(해당 원재료가 있는 경우)",
    "sjph_manual": "SJPH 매뉴얼 — 5요소(경영약속·원재료·공정·제품·모니터링) 포함",
    "logistics_scope": "취급 화물·서비스 범위 — 어떤 제품군(식품·의약·화장품)을 penyimpanan·pengemasan·pendistribusian 중 어느 서비스로 다루는지",
    "warehouse_layout": "창고 배치도 — 할랄/비할랄 구획(Halal Zone)과 동선 분리가 드러나야 함",
    "vehicle_list": "차량·컨테이너 목록 — 등록번호·유형. 인증 대상은 서비스이나 관리대상 증빙",
    "cleaning_sop": "세척 SOP — 차량·컨테이너·창고 세척 절차. 이전 화물이 비할랄일 때 Sertu 세정 포함",
}

# 서류 요건 설명 — 인니 신청기업·심사자가 읽는 문구다. 한국어만 두면 인니어 화면에
# 그대로 뜬다(실측: 사전심사 화면 잔여 한글의 최대 덩어리).
# 서류 요건 설명 — 사전에서 꺼낸다(표를 코드에 두 벌 두지 않는다).
from .domain_dict import code_labels as _code_labels   # noqa: E402

DOC_REQUIREMENT_L10N = {lg: _code_labels("DOC_REQUIREMENT", "doc_requirement", lg)
                        for lg in ("en", "id")}

# 면제·갈음 사유 — 왜 이 서류를 안 내도 되는지. 심사자가 읽고 판단하는 문장이다.
# 면제·갈음 사유 — 사전에서(한국어 문구가 키다).
def _exempt_l10n():
    from .domain_dict import text as _t
    kos = [v for tbl in list(DOC_NOT_APPLICABLE.values()) + [
        {k: x["note"] for k, x in v.items()} for v in DOC_ALT_SATISFY.values()]
        for v in [tbl] for v in tbl.values()]
    kos += [v for tbl in DOC_JURISDICTION.values() for v in tbl.values()]
    return {ko: {lg: _t(ko, lg) for lg in ("en", "id")} for ko in set(kos)}


_EXEMPT_L10N = _exempt_l10n()

def requirement_text(doc_type, lang="ko"):
    """서류 요건 설명 — 해당 언어가 없으면 한국어로 물러선다(빈칸보다 낫다)."""
    tbl = DOC_REQUIREMENT_L10N.get((lang or "ko").lower())
    return (tbl or {}).get(doc_type) or DOC_REQUIREMENT.get(doc_type, "")


def exempt_text(ko, lang="ko"):
    """면제·갈음 사유 문구를 해당 언어로."""
    if (lang or "ko").lower() == "ko":
        return ko
    return (_EXEMPT_L10N.get(ko) or {}).get((lang or "").lower(), ko)

_IMG = ("png", "jpg", "jpeg", "bmp", "tiff", "tif", "webp")


def _zip_name(zi):
    """한국어 Windows ZIP은 파일명이 CP949인데 UTF-8 플래그(0x800)가 없으면
    zipfile이 CP437로 디코드해 깨진다(┴╓..). CP437로 되돌려 CP949로 재디코드."""
    name = zi.filename
    if not (zi.flag_bits & 0x800):
        for enc in ("cp949", "euc-kr", "utf-8"):
            try:
                return name.encode("cp437").decode(enc)
            except Exception:
                continue
    return name


class ArchiveError(Exception):
    """압축을 열 수 없음 — 호출부가 사용자에게 사유를 그대로 보여준다."""


# 형식은 확장자가 아니라 매직바이트로 가린다. 이름만 .zip 으로 바꿔 보내는 경우가 있다.
_MAGIC = [(b"PK\x03\x04", "zip"), (b"PK\x05\x06", "zip"), (b"PK\x07\x08", "zip"),
          (b"Rar!\x1a\x07", "rar"), (b"7z\xbc\xaf\x27\x1c", "7z")]


def archive_kind(data):
    for sig, kind in _MAGIC:
        if data[:len(sig)] == sig:
            return kind
    return None


def _extract_with_tool(data, suffix):
    """bsdtar(macOS 기본 내장 libarchive)로 RAR·7z 를 푼다. 없으면 unar 로 폴백.

    파이썬 rarfile 패키지도 결국 외부 unrar 바이너리를 요구한다. bsdtar 는 이미 깔려
    있으므로 새 의존성 없이 같은 일을 한다."""
    import shutil
    import subprocess
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        src = os.path.join(td, "archive" + suffix)
        dst = os.path.join(td, "out")
        os.makedirs(dst)
        with open(src, "wb") as f:
            f.write(data)
        if shutil.which("bsdtar"):
            cmd = ["bsdtar", "-xf", src, "-C", dst]
        elif shutil.which("unar"):
            cmd = ["unar", "-quiet", "-force-overwrite", "-output-directory", dst, src]
        else:
            raise ArchiveError("이 서버에 %s 해제 도구가 없습니다 (bsdtar·unar 미설치)"
                               % suffix.lstrip("."))
        r = subprocess.run(cmd, capture_output=True, timeout=180)
        if r.returncode != 0:
            raise ArchiveError("압축을 풀지 못했습니다: %s"
                               % (r.stderr.decode("utf-8", "ignore")[:200] or "unknown"))
        out = []
        for root, _dirs, names in os.walk(dst):
            for n in sorted(names):
                if n.startswith("."):
                    continue
                p = os.path.join(root, n)
                rel = os.path.relpath(p, dst)
                if "__MACOSX" in rel:
                    continue
                try:
                    with open(p, "rb") as f:
                        out.append((rel, f.read()))
                except Exception:  # noqa: BLE001
                    pass
        return out


def extract_zip(data):
    """압축 묶음 → [(파일명, 바이트)]. ZIP·RAR·7z 를 모두 받는다.

    함수명은 zip 시절 그대로 두되(호출부가 여럿) 실제로는 형식을 가려서 연다.
    인도네시아 업체는 RAR로 보내는 일이 흔하다 — 예전에는 BadZipFile 로 500이 났고
    화면에는 '인테이크 완료'라고 떠서 서류 0건인 채 넘어갔다(실측)."""
    kind = archive_kind(data)
    if kind in ("rar", "7z"):
        return _extract_with_tool(data, "." + kind)
    if kind is None:
        raise ArchiveError("압축 파일이 아닙니다 (ZIP·RAR·7z 만 지원합니다)")
    out = []
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e:
        raise ArchiveError("ZIP 파일이 손상되었습니다: %s" % e) from e
    with z:
        for zi in z.infolist():
            if zi.is_dir():
                continue
            name = _zip_name(zi)
            base = os.path.basename(name)
            if not base or base.startswith(".") or "__MACOSX" in name:
                continue
            try:
                out.append((name, z.read(zi)))
            except Exception:  # noqa: BLE001
                pass
    return out


def _decode_text(data):
    """OS/로케일 독립 텍스트 디코드 — UTF-8→CP949→EUC-KR 순 시도(한국 Windows 파일 대응)."""
    for enc in ("utf-8", "cp949", "euc-kr", "latin-1"):
        try:
            return data.decode(enc)
        except Exception:
            continue
    return data.decode("utf-8", "ignore")


_OCR_MIN_CONF = float(os.environ.get("GLHAC_OCR_MIN_CONF", "0.5"))
_OCR_DPI = int(os.environ.get("GLHAC_OCR_DPI", "140"))   # 깨끗한 스캔은 140이 충분(측정), 저품질만 env 상향


def _ocr_bytes(data, ext, sink=None):
    """OS 독립 임시파일 OCR. confidence 낮은(오인식) 라인 제거로 품질↑.
    (LLM 한글교정은 qwen2.5 테스트 결과 성분명 환각으로 더 악화 → 미채택, 할랄 안전)

    sink 를 주면 인식 라인(좌표 포함)을 그대로 담아준다 — 표 사진의 행 복원은 좌표가
    있어야 하고, 없으면 나중에 같은 사진을 또 OCR하게 된다."""
    import tempfile
    fd, path = tempfile.mkstemp(suffix="." + ext)
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(data)
        lines = ai_local.ocr_image(path).get("lines", [])
        if sink is not None:
            sink.extend(lines[:4000])     # 병적으로 큰 스캔본 방어

        kept = [l["text"] for l in lines if l.get("confidence", 1) >= _OCR_MIN_CONF]
        if not kept and lines:            # 전부 저confidence면 폴백(빈 텍스트 방지)
            kept = [l["text"] for l in lines]
        return " ".join(kept)
    except Exception:  # noqa: BLE001
        return ""
    finally:
        try:
            os.remove(path)
        except OSError:
            pass


# 전성분표(원료명|INS No.|용도 반복 구조) 전용 구조 인식 파서.
# 근본 해결: 2차원 표를 " ".join으로 1차원화하면 열 의미가 소실돼(원료명·용도·제품명 뒤섞임,
# 인접 셀 병합) LLM이 오분류한다. 표 구조 그대로 제품명(시트 상단)·원재료(원료명 열)를 결정론적 추출.
_ING_HDR_RE = re.compile(r"원료\s*명|ingredient\s*name")
_ING_SKIP_RE = re.compile(r"^(원료\s*명|INS|용도|1차|2차|3차|ingredient|function|no\.?)$", re.I)

# 인니어 원재료 목록(Daftar Bahan): 헤더가 'Nama'/'Nama Bahan' 이라 위 한국어 규칙에
# 걸리지 않았다. 그 결과 136행짜리 표를 청크 LLM으로 다시 읽었고 — 87초를 쓰고 134개만
# 건졌다(실측). 표에 그대로 있는 값을 결정론적으로 읽는다.
# 헤더 셀은 '한 줄 = 한 언어'로 겹쳐 쓰인다("Nama Bahan (Material Name)\n재료명").
# 예전 규칙은 셀 전체가 정확히 'Nama'/'Nama Bahan' 일 때만 걸려서, GL-HAC 정식 서식
# Form.5 처럼 괄호 병기·한국어 병기가 붙은 표를 통째로 놓쳤다 — 그 표를 LLM 이 대신
# 읽었고 45건짜리 목록을 모델마다 39/45/54건으로 다르게 냈다(실측).
_ID_NAME_HDR_RE = re.compile(
    r"^\s*(nama(\s*bahan)?|material\s*name|ingredient\s*name|원료\s*명|재료\s*명)\b", re.I)
# 일련번호 열 — 이 열에 번호가 있는 행만 데이터로 본다(하위헤더·주석·합계행 배제).
_ID_NO_HDR_RE = re.compile(r"^\s*(no|번호)\.?\s*$", re.I)
# 이름 열 아래 언어 하위헤더(KOR/ENG). 원재료명으로 새어들면 목록 맨 앞에 박힌다.
_ID_SUB_HDR_RE = re.compile(r"^(kor|eng|ko|en|한글|영문|국문)$", re.I)


def _hdr_first_line(cell):
    """헤더 셀의 첫 줄 — 병기된 다른 언어 줄에 규칙이 헛걸리지 않게."""
    return str(cell).splitlines()[0].strip() if cell is not None else ""
# 원재료표임을 뒷받침하는 이웃 헤더 — 'Nama' 한 단어만으로 판단하면 아무 표나 걸린다.
_ID_ATTR_HDR_RE = re.compile(
    r"jenis\s*bahan|produsen|negara|supplier|pemasok|lembaga|sertifikat|no\.?\s*sert", re.I)
# 제품×원재료 매트릭스 표시값 — 이런 열이 이어지면 목록이 아니라 매트릭스다.
_MATRIX_MARK_RE = re.compile(r"^[vV✓xX✔○●\-–—]$")


def _looks_like_matrix(rows, hdr, name_col):
    """제품×원재료 매트릭스인가 — 목록과 같은 헤더를 쓰지만 성격이 다르다.

    'Bahan vs Produk matriks.xlsx' 는 'Nama Bahan' 헤더를 쓰면서 오른쪽이 전부 제품 열이고
    칸은 V 표시다. 이걸 원재료표로 읽으면 doc_type 이 product_list 에서 material_list 로
    뒤집혀 제품 목록이 사라진다. 표시값 열이 3개 이상이면 매트릭스로 본다."""
    body = rows[hdr + 1:hdr + 12]
    width = max((len(r) for r in body), default=0)
    marks = 0
    for j in range(name_col + 1, min(name_col + 12, width)):
        vals = [str(r[j]).strip() for r in body
                if j < len(r) and r[j] is not None and str(r[j]).strip()]
        if vals and all(_MATRIX_MARK_RE.match(v) for v in vals):
            marks += 1
    return marks >= 3


MATRIX_MARK = "[구조:제품×원재료매트릭스]"


def _id_material_rows(rows, seen=None):
    """인니어 원재료 목록 시트면 원재료명 리스트, 아니면 None.

    매트릭스로 판정하면 seen['matrix']=True 로 알린다. 그 사실을 남기지 않으면
    구조 파서는 '목록 아님'이라고 옳게 판단해 놓고, 같은 본문을 LLM 이 읽어 목록을
    지어낸다(실측: 같은 파일에서 로컬 LLM 이 제품 35건·원재료 21건을 뽑았고
    공급자·실행마다 값이 달랐다)."""
    for i, row in enumerate(rows[:8]):
        cols = [j for j, c in enumerate(row)
                if c and _ID_NAME_HDR_RE.match(_hdr_first_line(c))]
        if not cols:
            continue
        name_col = cols[0]
        # 매트릭스 검사를 먼저 한다. 이웃 헤더(Jenis/Produsen…) 개수로 먼저 걸러내면
        # 오른쪽이 전부 '제품 열'인 매트릭스는 attrs=0 이라 검사에 닿지도 못한다
        # (실측: 'Bahan vs Produk matriks.xlsx' 가 그랬고, 그 사이 LLM 이 목록을 지어냈다).
        if _looks_like_matrix(rows, i, name_col):
            if seen is not None:
                seen["matrix"] = True
            return None
        attrs = sum(1 for c in row if c and _ID_ATTR_HDR_RE.search(str(c)))
        if attrs < 2:            # 이웃 헤더가 없으면 원재료표라고 볼 근거가 없다
            continue
        # 일련번호 열이 있으면 그 열이 표의 경계다 — 번호 있는 행만 데이터로 본다.
        # 이름 열은 하위 언어열(KOR/ENG)로 쪼개져 있을 수 있어, 다음 헤더 열 직전까지를
        # 한 묶음으로 보고 그 안에서 처음 채워진 값을 쓴다(KOR 비면 ENG).
        no_col = next((j for j, c in enumerate(row)
                       if c and _ID_NO_HDR_RE.match(_hdr_first_line(c))), None)
        nxt = next((j for j, c in enumerate(row)
                    if j > name_col and c and str(c).strip()), None)
        span = range(name_col, nxt if nxt is not None else name_col + 1)
        out = []
        for row2 in rows[i + 1:]:
            if no_col is not None:
                seq = row2[no_col] if no_col < len(row2) else None
                seq = "" if seq is None else str(seq).strip().rstrip(".")
                if not seq.isdigit():
                    continue
                v = next((str(row2[j]).strip() for j in span
                          if j < len(row2) and row2[j] is not None
                          and str(row2[j]).strip()), "")
            else:
                if name_col >= len(row2) or not row2[name_col]:
                    continue
                v = str(row2[name_col]).strip()
            if (v and not _ING_SKIP_RE.match(v) and not _ID_NAME_HDR_RE.match(v)
                    and not _ID_SUB_HDR_RE.match(v) and v not in out):
                out.append(v)
        return out or None
    return None


def _xlsx_ingredient_table(wb):
    """전성분표 구조가 감지되면 (products, materials) 반환, 아니면 None.
    - 제품명 = 각 시트 1행의 첫 비어있지 않은 셀(완제품 1개/시트)
    - 원재료 = '원료명' 헤더가 있는 열들의 데이터 값만(INS No.·용도 열 제외)"""
    products, materials = [], []
    detected = False
    matrix_seen = False
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if not rows:
            continue
        # 헤더행: '원료명' 헤더를 포함하는 첫 행 → 원료명 열 인덱스 수집
        name_cols, hdr = [], None
        for i, row in enumerate(rows[:8]):
            cols = [j for j, c in enumerate(row) if c and _ING_HDR_RE.search(str(c))]
            if cols:
                name_cols, hdr = cols, i
                break
        if hdr is None:
            # 인니어 원재료 목록(Nama + Jenis Bahan/Produsen/…). 제품명은 없다 —
            # 시트 1행은 표 제목('Daftar Bahan Halal …')이라 제품으로 쓰면 안 된다.
            _seen = {}
            id_mats = _id_material_rows(rows, _seen)
            if _seen.get("matrix"):
                matrix_seen = True
            if id_mats:
                detected = True
                for v in id_mats:
                    if v not in materials:
                        materials.append(v)
            continue   # 이 시트는 (한국어) 전성분표 구조 아님
        detected = True
        prod = next((str(c).strip() for c in rows[0] if c and str(c).strip()), None)
        if prod and prod not in products:
            products.append(prod)
        for row in rows[hdr + 1:]:
            for j in name_cols:
                if j < len(row) and row[j]:
                    v = str(row[j]).strip()
                    if v and not _ING_SKIP_RE.match(v) and v not in materials:
                        materials.append(v)
    if matrix_seen and not detected:
        # 매트릭스만 있고 목록은 없다 — '아무것도 못 찾음'과 구분해 돌려준다.
        return "matrix"
    if not detected:
        return None
    return products, materials


def parse_file(name, data, dpi=None, ocr_sink=None):
    """확장자별 텍스트 추출 — OS 독립. 이미지/스캔PDF=OCR, PDF=fitz, docx/xlsx/txt.
    dpi 지정 시 스캔 렌더 해상도 오버라이드(고해상도 재처리용, 기본 _OCR_DPI).

    ocr_sink 를 주면 OCR 라인(좌표 포함)을 담아준다 — 재-OCR을 막는 캐시용. 단일 이미지에만
    채운다. 스캔 PDF는 페이지마다 좌표계가 새로 시작해 여러 페이지를 한 목록에 담으면 서로
    다른 페이지의 행이 같은 y 값으로 겹친다 — 표 행 복원이 조용히 틀어지느니 캐시하지 않는다."""
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    _dpi = int(dpi) if dpi else _OCR_DPI
    _dpi = max(72, min(600, _dpi))   # 안전 범위
    try:
        if ext in _IMG:
            return _ocr_bytes(data, ext, ocr_sink)
        if ext == "pdf":
            import fitz
            doc = fitz.open(stream=data, filetype="pdf")
            txt = ""
            for pg in list(doc)[:8]:
                t = pg.get_text()
                if len(t.strip()) < 20:  # 스캔본 → PNG 렌더 후 OCR (파일경로 없이 bytes)
                    t = _ocr_bytes(pg.get_pixmap(dpi=_dpi).tobytes("png"), "png")
                txt += t + "\n"
            return txt
        if ext in ("txt", "csv"):
            return _decode_text(data)
        if ext == "docx":
            import docx
            d = docx.Document(io.BytesIO(data))
            return "\n".join(pp.text for pp in d.paragraphs)
        if ext == "xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            # 근본 해결: 전성분표 구조면 표를 구조 인식으로 파싱해 명시적 마커로 넘긴다(LLM 우회).
            ing = _xlsx_ingredient_table(wb)
            if ing == "matrix":
                # 제품×원재료 매트릭스 — 목록 문서가 아니다. 그 사실을 본문에 남겨
                # LLM 이 이 표에서 목록을 지어내도 집계가 받지 않게 한다.
                ing = None
                matrix_note = MATRIX_MARK
            else:
                matrix_note = ""
            if ing:
                wb.close()
                prods, mats = ing
                out = ["[전성분표]"]
                for p in prods:
                    out.append("[제품/시트: %s]" % p)
                out.append("[원재료목록]")
                out.extend(mats)
                return "\n".join(out)[:60000]
            # 일반 xlsx(전성분표 아님): 기존 flatten
            out = [matrix_note] if matrix_note else []
            for ws in wb.worksheets:
                title = (ws.title or "").strip()
                if title and title.lower() not in ("sheet", "sheet1", "sheet2", "sheet3", "시트1"):
                    out.append("[시트/제품명: %s]" % title)
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        out.append(" ".join(cells))
            wb.close()
            return "\n".join(out)[:20000]
    except Exception:  # noqa: BLE001
        return ""
    return ""


_CLASSIFY_SYS = (
    "당신은 할랄 인증 서류 분류기입니다. 파일명과 본문 발췌를 보고 doc_type을 분류하고 핵심 필드를 추출하세요. "
    "doc_type은 반드시 다음 중 하나: " + ", ".join(DOC_TYPES) + ". "
    "address는 회사(사업자) 주소, city/province/country/zip은 회사 주소의 도시/도(주,state/province)/국가/우편번호. "
    "factory_address는 공장 주소, factory_city/factory_province/factory_country/factory_zip은 공장 주소의 도시/도(주)/국가/우편번호. "
    "factory_reg_no는 공장등록번호(Plant Registration No. — 공장 고유번호, 예: G20026130). "
    "responsible_person은 대표자/책임자 이름. "
    "phone은 회사 전화번호, factory_phone은 공장 전화번호(Tel). "
    "business_type은 사업의 종류/업태(예: 제조업, 도소매업). "
    "employee_count는 총 직원 수(숫자만, 남녀 합계). "
    "establishment_date는 개업일/설립일(YYYY-MM-DD). "
    "corporate_reg_no는 법인등록번호(Corporate Registration No., 형식 123456-1234567 — 공장등록번호 factory_reg_no와 다른 번호이니 혼동하지 말 것). "
    "주소(address/factory_address)는 도로명부터 도·시·국가까지 전체를 그대로 넣으세요(중간에서 자르지 말 것). "
    "process_steps는 공정흐름도의 공정 단계 순서 목록(원료입고→배합→가열→포장 등). "
    "product_names는 완제품 이름 목록 — 제품목록 문서뿐 아니라 전성분표/원재료 문서라도 "
    "시트명(예: '[시트/제품명: ...]')·표 제목·문서 상단에 완제품명이 있으면 product_names에 넣으세요. "
    "material_names는 원재료(성분) 이름 목록(완제품이 아님). "
    "has_commitment/has_materials/has_process/has_product/has_monitoring는 SJPH 매뉴얼에 "
    "해당 5요소(약속·원재료·공정·제품·모니터링)가 포함되면 true. 없으면 null/false/[]. "
    '반드시 JSON으로만: {"doc_type":"...","confidence":0.0,'
    '"fields":{"company_name":null,"nib":null,"address":null,"city":null,"province":null,"country":null,"zip":null,'
    '"factory_address":null,"factory_city":null,"factory_province":null,"factory_country":null,"factory_zip":null,'
    '"responsible_person":null,"factory_reg_no":null,"phone":null,"factory_phone":null,'
    '"business_type":null,"employee_count":null,"establishment_date":null,"corporate_reg_no":null,'
    '"product_names":[],"cert_no":null,'
    '"issuer":null,"expiry_date":null,"material_names":[],"process_steps":[],'
    '"has_commitment":false,"has_materials":false,"has_process":false,'
    '"has_product":false,"has_monitoring":false}}'
)


# 파일명 규칙 — LLM 오분류 교정. 위에서부터 먼저 매칭되는 규칙이 최종 doc_type을 결정한다.
# 파일명이 명백한 경우(소개서·제안서·수입서류·품질인증 등)는 LLM 판단을 무시하고 규칙을 신뢰한다.
_NAME_RULES = [
    # (정규식, doc_type, 사유) — 순서 중요(구체적인 것 먼저)
    (r"bahan\s*vs\s*produk|produk\s*matriks|matriks\s*produk|product.{0,3}matri|matri.{0,3}produk",
     "product_list", "제품×원재료 매트릭스(제품 카탈로그 — 제품명 컬럼이 정본 제품 목록)"),
    (r"\bcatatan\b|기록부|생산\s*일지|작업\s*일지|monitoring\s*record|rekaman",
     "other", "운영·모니터링 기록(구매·생산·재고·검사 기록 — 제품/원재료 카탈로그 아님)"),
    (r"management[\s_]*review|tinjauan[\s_]*manajemen|관리\s*검토|경영\s*검토|"
     r"halal[\s_]*policy|kebijakan[\s_]*halal|할랄\s*정책|"
     r"internal[\s_]*audit|audit[\s_]*internal|내부\s*심사|"
     r"halal[\s_]*team|tim[\s_]*halal|할랄\s*팀|"
     r"halal[\s_]*manual|manual[\s_]*halal|\bsjph\b",
     "sjph_manual", "SJPH 구성문서(할랄정책·할랄팀·내부심사·경영검토·매뉴얼 — 사업자/공급사 아님)"),
    (r"제조\s*공정\s*도|공정\s*흐름|process\s*flow|flow\s*chart", "process_flow", "공정도"),
    (r"fssc|haccp|\biso\b|\bgmp\b|22000|식품안전|유기취급|organic|kosher", "quality_cert", "품질/식품안전 인증(HACCP·FSSC·ISO·GMP — 할랄 아님)"),
    (r"소개서|회사\s*소개|company\s*profile|제안서|proposal|접수\s*양식|고객\s*접수|데이터\s*양식|intake\s*form|application\s*form", "other", "소개서·제안서·양식(등록증 아님)"),
    (r"원산지|수입\s*서류|country\s*of\s*origin|\bcoo\b|certificate\s*of\s*origin|선언서|declaration|확인서|설명서", "supplier_declaration", "원산지·수입·선언서(공급사 선언)"),
    (r"성적서|시험\s*성적|\bcoa\b|\bmsds\b|성분\s*명세|성분\s*분석|분석\s*성적", "coa_msds", "성적서·성분명세"),
    (r"사업자\s*등록|사업자등록증|business\s*(registration|license)|\bnib\b|법인\s*등기|등록증명원", "nib_business_license", "사업자등록"),
    (r"공장\s*등록|공장등록증|factory\s*registration|manufactur.*licen", "factory_registration", "공장등록"),
    (r"할랄\s*인증|halal\s*cert", "halal_certificate", "할랄 인증"),
]


# 서류 제목으로 볼 만한 줄 — 표지·시트명·머리글. 본문 아무 데나 찾으면 원재료표 안의
# 한 단어에 걸려 유형이 뒤집힌다. 앞부분의 짧은 줄만 본다.
_TITLE_MAX_LINES = 12
_TITLE_MAX_LEN = 90
_SHEET_MARK_RE = re.compile(r"^\[시트/제품명:\s*(.+?)\]\s*$")


def _title_lines(text):
    """문서가 스스로 밝힌 제목 후보. 시트명 마커는 전부, 본문은 앞쪽 짧은 줄만."""
    if not text:
        return []
    out, plain = [], 0
    for ln in text.splitlines():
        ln = ln.strip()
        if not ln:
            continue
        m = _SHEET_MARK_RE.match(ln)
        if m:
            out.append(m.group(1).strip())
            continue
        if plain >= _TITLE_MAX_LINES:
            continue
        plain += 1
        if len(ln) <= _TITLE_MAX_LEN:
            out.append(ln)
    return out


def refine_doctype_reason(name, llm_type, text=None):
    """(doc_type, 규칙근거) — 파일명 규칙 → 사전(파일명) → 사전(제목줄) → LLM 판단 순.

    사전 폴백을 둔 이유: 인니어 원문 파일명('Diagram alir proses produksi',
    'Catatan pembelian barang')은 한국어·영어로 짜인 _NAME_RULES에 걸리지 않는다.
    사전은 세 언어 표면형을 한 표준 키로 모으므로 규칙을 언어마다 늘리지 않아도 된다.

    제목줄까지 보는 이유: 파일명은 고객이 마음대로 바꾼다. 'Form 9(피치×샷).xlsx' 는
    사전이 아무것도 모르지만, 그 안의 시트명은 'Form.9 재료 보관 기록' 이고 사전은
    그것을 운영 기록물로 정확히 안다. 파일명만 보던 동안 이 서류는 원재료 목록으로
    분류돼 보관 기록의 입출고 행에서 원재료 42~74건이 지어졌다(모델마다 달랐다)."""
    n = (name or "").lower()
    from .domain_dict import doc_type_of, evidence_key_of, lookup
    base = re.sub(r"\.[a-z0-9]{2,5}$", "", (name or "").strip())
    for pat, dt, why in _NAME_RULES:
        if re.search(pat, n):
            # 규칙이 잡았더라도 사전이 증빙 항목을 알고 있으면 근거에 함께 남긴다
            # (인니 실무 기록물은 doc_type은 other여도 SJPH 증빙 항목에 붙는다).
            ev = evidence_key_of(base)
            return dt, (why + " · 증빙 " + ev) if ev else why
    dt = doc_type_of(base)
    if dt:
        key = lookup(base, axis="DOC")
        ev = evidence_key_of(base)
        return dt, "도메인사전 %s%s" % (key, ("·증빙 " + ev) if ev else "")
    for title in _title_lines(text):
        dt = doc_type_of(title)
        if dt:
            key = lookup(title, axis="DOC")
            ev = evidence_key_of(title)
            return dt, "제목줄 '%s' → 도메인사전 %s%s" % (
                title[:40], key, ("·증빙 " + ev) if ev else "")
    return llm_type, None


def _refine_doctype(name, llm_type, text=None):
    return refine_doctype_reason(name, llm_type, text)[0]


# 리스트형 문서(원재료·제품·공정·성분)는 전체 본문에서 항목을 빠짐없이 추출한다.
# classify(2000자)·parse_typed(2500자) 절단으로 목록 뒤쪽이 통째로 누락되던 문제
# (예: Daftar bahan.xlsx 원재료 136개 → 18개만 추출)를 해결한다.
_LIST_FIELDS = {
    "material_list": ("material_names", "원재료명"),
    # product_list(제품×원재료 매트릭스)는 헤더행이 정식 제품목록이므로 기본 분류값을 신뢰한다.
    # 전체본문 list-completion을 돌리면 매트릭스의 원재료 셀까지 제품명으로 과다추출됨 → 제외.
    "process_flow": ("process_steps", "공정 단계"),
    "product_label": ("ingredients", "성분명"),
}


def _norm_key(s):
    """제품/원재료 중복 판정 키 — 대소문자·구두점·공백 무시(한글 등 유니코드 문자 보존).
    [^a-z0-9]로 지우면 한글 제품명이 빈 키가 돼 통째로 유실되므로 \\W_ 만 제거한다."""
    return re.sub(r"[\W_]", "", (s or "").lower())


def _chunk_text(text, size=4500):
    """줄 단위로 ~size자 청크 분할(항목이 중간에서 잘리지 않게)."""
    chunks, buf, n = [], [], 0
    for line in text.splitlines():
        ln = len(line) + 1
        if n + ln > size and buf:
            chunks.append("\n".join(buf)); buf, n = [], 0
        buf.append(line); n += ln
    if buf:
        chunks.append("\n".join(buf))
    return chunks or [text]


def _extract_list_complete(filename, text, label, max_chunks=12, cap=1000):
    """전체 본문을 청크로 나눠 목록 항목을 완전 추출·합집합(대소문자 무시 dedup)."""
    sys = ('문서에서 %s 목록을 빠짐없이 추출하세요. 표·번호목록의 모든 행을 포함하고 '
           '헤더/합계/빈칸은 제외. 반드시 JSON으로만: {"names":[]}' % label)
    seen, out = set(), []
    for ch in _chunk_text(text)[:max_chunks]:
        r = ai_local.llm_json(sys, "파일명: %s\n본문:\n%s" % (filename, ch))
        for nm in (r.get("names") if isinstance(r, dict) else None) or []:
            if not isinstance(nm, str):
                continue
            nm = nm.strip()
            k = _norm_key(nm)
            if nm and k and k not in seen:
                seen.add(k); out.append(nm)
                if len(out) >= cap:
                    return out
    return out


def _infer_country_zip(addr):
    """주소 문자열에서 국가/우편번호 추론(LLM이 분리 못했을 때 보정)."""
    if not addr:
        return None, None
    zipc = None
    m = re.search(r'\b(\d{5})\b', addr)
    if m:
        zipc = m.group(1)
    low = addr.lower()
    country = None
    if any(x in low for x in ("indonesia", "dki", "provinsi", "kota adm", "jakarta",
                              "kabupaten", "kelurahan", "kecamatan")):
        country = "Indonesia"
    elif "korea" in low:
        country = "Republic of Korea"
    return country, zipc


_ZIP_CACHE = {}
_ZIP_LOOKUP = os.environ.get("GLHAC_ZIP_LOOKUP", "1") == "1"


def _lookup_zip(address):
    """문서에 우편번호가 없을 때 주소로 우편번호를 조회(Nominatim/OpenStreetMap 지오코딩, 전세계).
    무료·키 불필요. 실패/미발견 시 None. 캐시. 한국 상세 도로명은 정확도가 낮을 수 있음."""
    if not _ZIP_LOOKUP or not address:
        return None
    a = str(address).strip()
    if not a or a in _ZIP_CACHE:
        return _ZIP_CACHE.get(a)
    z = None
    try:
        import httpx
        r = httpx.get("https://nominatim.openstreetmap.org/search",
                      params={"q": a, "format": "json", "addressdetails": 1, "limit": 1},
                      headers={"User-Agent": "GLHAC-Halal-Cert/1.0"}, timeout=6)
        d = r.json()
        if d:
            z = (d[0].get("address") or {}).get("postcode")
    except Exception:  # noqa: BLE001 — 네트워크 장애 시 조용히 폴백(추출 차단 없음)
        z = None
    _ZIP_CACHE[a] = z
    return z


def _complete_address(addr, city, province, country):
    """LLM이 주소 뒷부분(시·도·국가)을 절단한 경우 도시→도(주)→국가 순으로 이어붙여 완성.
    이미 포함돼 있으면 그대로 둔다(중복 방지). 표기 순서(…시, 도, 국가)를 지킨다."""
    if not addr:
        return addr
    out = str(addr).rstrip(" .,")
    low = out.lower()
    for part in (city, province, country):   # 도시 → 도(주) → 국가 순
        if part and str(part).strip() and str(part).lower() not in low:
            out += ", " + str(part).strip()
            low = out.lower()
    return out


def _enrich_address(fields):
    """LLM이 country/zip을 비웠으면 주소에서 보정(도시는 LLM 프롬프트에 의존)."""
    co, zp = _infer_country_zip(fields.get("address"))
    if co and not fields.get("country"):
        fields["country"] = co
    if zp and not fields.get("zip"):
        fields["zip"] = zp
    fco, fzp = _infer_country_zip(fields.get("factory_address"))
    if fco and not fields.get("factory_country"):
        fields["factory_country"] = fco
    if fzp and not fields.get("factory_zip"):
        fields["factory_zip"] = fzp


def _parse_ingredient_markers(text):
    """parse_file이 전성분표에 심은 구조 마커에서 (products, materials) 추출. 없으면 None.
    구조 파서 결과이므로 LLM 목록추출을 대체(오분류·셀병합·비결정성 제거)."""
    if "[전성분표]" not in text:
        return None
    prods = [m.strip() for m in re.findall(r'\[제품/시트:\s*([^\]]+)\]', text) if m.strip()]
    mats = []
    if "[원재료목록]" in text:
        body = text.split("[원재료목록]", 1)[1]
        mats = [ln.strip() for ln in body.splitlines() if ln.strip() and not ln.startswith("[")]
    return list(dict.fromkeys(prods)), list(dict.fromkeys(mats))


# 시험성적서(CoA/MSDS) 정량 측정값 추출 — 정규식(결정적). param_key는 main.QUANT_CRITERIA와 일치.
_QUANT_PATTERNS = {
    "ethanol_pct":  r"(?:ethanol|알코올|에탄올|alcohol)[^\d%]{0,25}?([\d.]+)\s*%",
    "lead_ppm":     r"(?:lead|\bPb\b|납)[^\d]{0,25}?([\d.]+)\s*(?:ppm|mg/kg)",
    "cadmium_ppm":  r"(?:cadmium|\bCd\b|카드뮴)[^\d]{0,25}?([\d.]+)\s*(?:ppm|mg/kg)",
    "mercury_ppm":  r"(?:mercury|\bHg\b|수은)[^\d]{0,25}?([\d.]+)\s*(?:ppm|mg/kg)",
    "arsenic_ppm":  r"(?:arsenic|\bAs\b|비소)[^\d]{0,25}?([\d.]+)\s*(?:ppm|mg/kg)",
}


def extract_measurements(text):
    """시험성적서 본문에서 정량 파라미터 측정값을 추출(결정적 정규식). {param_key: value(float)}."""
    if not text:
        return {}
    out = {}
    for k, pat in _QUANT_PATTERNS.items():
        m = re.search(pat, text, re.I)
        if m:
            try:
                out[k] = float(m.group(1))
            except ValueError:
                pass
    # 돼지 DNA: 검출/불검출(정성)
    if re.search(r"(?:pork|porcine|돼지|돈지|babi)[^\n]{0,20}?(?:dna|pcr)", text, re.I):
        neg = re.search(r"(?:not\s*detect|non[- ]?detect|negative|불검출|음성|없음|none|absent)", text, re.I)
        out["pork_dna"] = 0.0 if neg else 1.0
    return out


# 파일명만으로 종류가 확정되고, LLM 이 뽑는 필드도 쓰지 않는 서류들.
# 이 경우 LLM 호출은 결과가 어차피 파일명 규칙으로 덮어써져 순수 낭비다(파일당 ~25초).
# 실측: 16건 인테이크 9분 30초 중 7분이 LLM 대기였고 11건은 파일명으로 이미 판별됐다.
_NAME_DECIDES = {"sjph_manual", "supplier_declaration", "process_flow",
                 "quality_cert", "origin_certificate", "halal_certificate"}


def _llm_failure(r):
    """LLM 호출이 실패했으면 {'llm_error': 사유}, 아니면 {}.

    'AI 없음' 배포(LLM_UNAVAILABLE)는 실패가 아니라 설계된 상태다 — 그 모드에서 모든
    서류를 확인 대기로 밀어 올리면 대기 목록이 의미를 잃는다."""
    if not isinstance(r, dict):
        return {"llm_error": "BAD_RESPONSE"}
    err = r.get("error")
    if not err or err == "LLM_UNAVAILABLE":
        return {}
    return {"llm_error": str(err)[:200]}


def _structural_material_result(ing):
    """구조 파서가 읽은 전성분표/원재료표 결과 — LLM 이 관여하지 않은 확정값.

    이 목록은 LLM 이 덮지 못한다. 같은 서류가 배포(없음/로컬/원격)마다 다른 원재료를
    내면 심사 결과가 갈린다."""
    prods, mats = ing
    fields = {}
    if prods:
        fields["product_names"] = prods     # 시트당 완제품 1개(A1)
    if mats:
        fields["material_names"] = mats     # 이름 열 값만
    _enrich_address(fields)
    srcs = {k: "structural" for k in ("product_names", "material_names") if fields.get(k)}
    return {"doc_type": "material_list", "confidence": 0.95, "fields": fields,
            "decided_by": "structural", "field_sources": srcs,
            "locked_fields": sorted(srcs)}


def classify(name, text):
    if not text.strip():
        # 본문이 비어도(이미지 PDF·OCR 미가동) 파일명이 아는 유형은 살린다.
        # 종전에는 그냥 other 로 떨어뜨려, 파일명 사전이 답을 알고 있는데도 버렸다 —
        # AI 없는 배포에서는 이 경로가 유일한 판정 수단이다.
        _dt, _why = refine_doctype_reason(name, None, text)
        if _dt and _dt != "other":
            return {"doc_type": _dt, "confidence": 0.6, "fields": {}, "empty": True,
                    "decided_by": "filename", "reason": _why}
        return {"doc_type": "other", "confidence": 0.0, "fields": {}, "empty": True}
    dt_by_name, why = refine_doctype_reason(name, None, text)
    if dt_by_name in _NAME_DECIDES:
        return {"doc_type": dt_by_name, "confidence": 0.9, "fields": {},
                "decided_by": "filename", "reason": why}
    # 사전이 '운영 기록물'(구매·생산·재고·검사 기록)로 확정한 문서는 LLM을 부르지 않는다.
    # 불러봐야 뽑은 값이 전부 버려진다 — other 는 _APPLICANT_DOCS·_PRODUCT_SRC·
    # _MATERIAL_SRC 어디에도 없어 aggregate_fields 가 회사정보·제품·원재료를 모두
    # 무시하고, _process_doc 도 다루지 않는다(실측: 기록 사진 6장에서 cert_no 는 전부
    # None, 나머지 필드는 집계 진입조차 못 함). 파일당 15~20초를 그렇게 썼다.
    # dt_by_name 이 'other' 인 경우는 사전이 아는 문서뿐이다 — doc_type_of 는 모르는
    # 이름·제목에 None 을 준다. 따라서 미지 문서의 LLM 판정을 뺏지 않는다.
    if dt_by_name == "other" and why:
        return {"doc_type": "other", "confidence": 0.9, "fields": {},
                "decided_by": "filename", "reason": why}
    # 구조 파서가 전성분표/원재료표를 이미 읽었으면 LLM 을 부르지 않는다.
    # 종전에는 부른 뒤 그 결과를 덮었다 — 목록은 지켜졌지만 서류당 10~90초를 버렸고,
    # 더 나쁜 건 그 호출이 실패하면 같은 서류가 배포·시점에 따라 다른 상태로 남았다는
    # 점이다(본문이 마커뿐이라 LLM 이 낼 수 있는 건 지어낸 주소 정도였다).
    ing = _parse_ingredient_markers(text)
    if ing is not None:
        return _structural_material_result(ing)
    r = ai_local.llm_json(_CLASSIFY_SYS, "파일명: %s\n본문 발췌:\n%s" % (name, text[:2000]))
    if not isinstance(r, dict) or "doc_type" not in r:
        # 호출이 실패한 것과 '모델이 보고도 못 찾은 것'은 다르다. 둘 다 other/0.0 으로
        # 떨어뜨리면 화면에는 똑같이 '분석 완료, 특이사항 없음'으로 보인다 — 실측:
        # 다른 서비스가 Ollama 를 점유한 동안 서류가 90초 타임아웃으로 줄줄이 빈 결과가
        # 됐는데 실패한 티가 어디에도 없었다. 실패는 실패라고 남겨 사람이 다시 돌린다.
        r = {"doc_type": "other", "confidence": 0.0, "fields": {},
             **_llm_failure(r)}
    r["doc_type"] = _refine_doctype(name, r.get("doc_type"), text)  # 파일명·제목줄 교정
    if r.get("doc_type") not in DOC_TYPES:
        r["doc_type"] = "other"
    r.setdefault("confidence", 0.0)
    r.setdefault("fields", {})
    # 리스트형 문서는 전체 본문에서 목록을 완전 추출(절단 2000자로는 뒤쪽 누락).
    lf = _LIST_FIELDS.get(r["doc_type"])
    if lf and len(text) > 2000:
        fk, label = lf
        complete = _extract_list_complete(name, text, label)
        if len(complete) > len(r["fields"].get(fk) or []):
            r["fields"][fk] = complete
    # 구조 파서가 '이건 목록 문서가 아니다'라고 판정했으면(제품×원재료 매트릭스)
    # LLM 이 뽑은 목록은 버린다. 구조 판정이 옳고 LLM 이 지어낸 것이다 —
    # 실측: 같은 매트릭스 파일에서 로컬 LLM 이 제품 35건, GPT 가 원재료 21건을 뽑았고
    # 실행마다 값이 달라 배포별로 신청서 내용이 갈렸다.
    if MATRIX_MARK in text:
        for _lk in ("material_names", "product_names"):
            r["fields"].pop(_lk, None)
        r["field_sources"] = {"material_names": "structural_reject",
                              "product_names": "structural_reject"}
        r["locked_fields"] = ["material_names", "product_names"]
        r["reason"] = "제품×원재료 매트릭스 — 목록 문서가 아님"
        _enrich_address(r["fields"])
        return r
    # 구조 파서가 안 걸린 목록은 LLM 이 낸 값이다. 출처를 남겨 두면 화면이
    # '사람이 확인할 값'으로 표시할 수 있고, 배포별 차이도 추적된다.
    r.setdefault("field_sources", {})
    for _lk in ("material_names", "product_names"):
        if r["fields"].get(_lk) and _lk not in r["field_sources"]:
            r["field_sources"][_lk] = "llm"
    r.setdefault("locked_fields", [])
    _enrich_address(r["fields"])   # 도시/국가/우편 보정
    # 시험성적서/성분명세/품질인증 → 정량 측정값 추출(자동 반영용)
    if r["doc_type"] in ("coa_msds", "quality_cert", "supplier_declaration"):
        _meas = extract_measurements(text)
        if _meas:
            r["fields"]["measurements"] = _meas
    return r


# 개별 업로드 컨텍스트 파싱 — doc_type별 추출 필드 스펙
_FIELD_SPEC = {
    "nib_business_license": ("회사명, NIB(사업자등록번호), 주소(도·시·국가 전체), 도시, 도(주), 국가, 우편번호, 회사 전화번호, 사업의 종류/업태, 개업일(YYYY-MM-DD), 법인등록번호",
                             '{"company_name":null,"nib":null,"address":null,"city":null,"province":null,"country":null,"zip":null,"phone":null,"business_type":null,"establishment_date":null,"corporate_reg_no":null}'),
    "factory_registration": ("공장등록번호, 공장 주소(도·시·국가 전체), 도시, 도(주), 국가, 우편번호, 공장 전화번호(Tel), 총 직원 수(숫자)",
                             '{"factory_reg_no":null,"factory_address":null,"factory_city":null,"factory_province":null,"factory_country":null,"factory_zip":null,"factory_phone":null,"employee_count":null}'),
    "halal_certificate": ("인증번호, 발급기관, 만료일, 대상(제품/원재료)",
                          '{"cert_no":null,"issuer":null,"expiry_date":null,"scope":null}'),
    "quality_cert": ("인증종류(HACCP/ISO/GMP/FSSC), 인증번호, 만료일",
                     '{"cert_type":null,"cert_no":null,"expiry_date":null}'),
    "consent": ("서명 여부, 서명일", '{"signed":false,"signed_date":null}'),
    "product_list": ("제품명 목록", '{"product_names":[]}'),
    "material_list": ("원재료명 목록. 표 제목·시트명(예: '[시트/제품명: ...]')·문서 상단에 제품명이 있으면 product_names에도 넣으세요(원재료가 아닌 완제품명)",
                      '{"material_names":[],"product_names":[]}'),
    "process_flow": ("공정 단계 순서 목록(원료입고→배합→가열→충전→포장 등)",
                     '{"process_steps":[]}'),
    "sjph_manual": ("SJPH 5요소 포함 여부(약속·원재료·공정·제품·모니터링)",
                    '{"has_commitment":false,"has_materials":false,"has_process":false,'
                    '"has_product":false,"has_monitoring":false}'),
    "product_label": ("성분 목록", '{"ingredients":[]}'),
}


def parse_typed(doc_type, filename, data):
    """업로드 칸(문구)에 맞는 doc_type으로 해당 필드만 파싱."""
    text = parse_file(filename, data)
    if not text.strip():
        return {"doc_type": doc_type, "fields": {}, "confidence": 0.0, "empty": True, "text_len": 0}
    # 근본: 전성분표 구조 마커가 있으면 구조 파서 결과를 신뢰(LLM 목록추출 우회).
    ing = _parse_ingredient_markers(text)
    if ing is not None:
        prods, mats = ing
        fields = {}
        if prods:
            fields["product_names"] = prods
        if mats:
            fields["material_names"] = mats
        return {"doc_type": "material_list", "fields": fields, "confidence": 0.95,
                "text_len": len(text), "excerpt": text[:300]}
    spec = _FIELD_SPEC.get(doc_type)
    if not spec:
        cl = classify(filename, text)
        cl["text_len"] = len(text)
        return cl
    desc, fmt = spec
    lf = _LIST_FIELDS.get(doc_type)
    if lf and len(text) > 2500:
        # 리스트형: 전체 본문에서 완전 추출(절단 방지)
        fk, label = lf
        fields, fail = {fk: _extract_list_complete(filename, text, label)}, {}
    else:
        sys = "문서에서 다음 필드만 추출하세요(없으면 null). 반드시 JSON으로만: " + fmt + " (대상: " + desc + ")"
        r = ai_local.llm_json(sys, "파일명: %s\n본문:\n%s" % (filename, text[:2500]))
        fail = _llm_failure(r)
        # 실패 응답({'error': ...})을 그대로 필드로 쓰면 error 라는 이름의 값이 저장되고
        # confidence 0.85 까지 붙는다 — 못 뽑았는데 잘 뽑은 것처럼 기록된다.
        fields = {} if fail or not isinstance(r, dict) else r
    _enrich_address(fields)   # 도시/국가/우편 보정
    out = {"doc_type": doc_type, "fields": fields,
           "confidence": 0.0 if fail else 0.85,
           "text_len": len(text), "excerpt": text[:300]}
    out.update(fail)
    return out


_APPLICANT_DOCS = ("nib_business_license", "factory_registration")
# 제품/원재료명을 신뢰할 카탈로그 문서 유형(기록·타목록의 오염 방지)
# material_list(전성분표)도 시트명 기반 완제품 추출 대상이므로 포함
_PRODUCT_SRC = {"product_list", "sjph_manual", "material_list"}
_MATERIAL_SRC = {"material_list", "coa_msds", "product_label"}


def aggregate_fields(docs):
    """분류 문서들의 추출 필드를 신청서용으로 집계.
    회사명·NIB·주소·책임자·공장등록번호는 신청기업 서류(사업자/공장등록증)에서만 취함(공급사 제외).

    같은 서류 묶음이면 LLM 공급자(없음·로컬·원격)와 무관하게 같은 결과를 낸다.
    구조 파서가 원재료·제품 목록을 낸 경우 그 필드는 결정적 값만 받는다."""
    agg = {"company_name": None, "company_name_ko": None, "company_name_en": None,
           "nib": None, "address": None, "factory_address": None,
           "city": None, "province": None, "country": None, "zip": None,
           "factory_city": None, "factory_province": None, "factory_country": None, "factory_zip": None,
           "responsible_person": None, "responsible_person_ko": None, "responsible_person_en": None,
           "phone": None, "factory_phone": None,
           "business_type": None, "employee_count": None,
           "establishment_date": None, "corporate_reg_no": None,
           "factory_reg_no": None, "products": [], "materials": [], "certificates": [],
           # LLM 만 뽑은 목록 — 자동 반영하지 않고 사람 확인을 기다린다
           "pending": []}
    _pk = set()   # 제품 중복 판정 키(대소문자 무시)
    _mk = set()   # 원재료 중복 판정 키
    _pend_pk, _pend_mk = set(), set()
    # 결정적 목록이 있으면 그 필드는 LLM 을 받지 않는다(공급자 무관 재현성).
    _struct_only = {"products": False, "materials": False}
    for _d in docs:
        _s = (_d.get("field_sources") or {})
        if _s.get("product_names") == "structural":
            _struct_only["products"] = True
        if _s.get("material_names") == "structural":
            _struct_only["materials"] = True
    for d in docs:
        f = d.get("fields") or {}
        applicant = d.get("doc_type") in _APPLICANT_DOCS or d.get("doc_type") is None
        if applicant:
            # 회사명 국문/영문 분리 — 한글 포함이면 국문(정식상호), 아니면 영문. 주 필드는 국문 우선.
            if f.get("company_name"):
                _cn = str(f["company_name"]).strip()
                if re.search(r'[가-힣]', _cn):
                    if not agg["company_name_ko"]:
                        agg["company_name_ko"] = _cn
                elif not agg["company_name_en"]:
                    agg["company_name_en"] = _cn
            if not agg["nib"] and f.get("nib"):
                agg["nib"] = f["nib"]
            # 회사 주소(NIB)와 공장 주소(공장등록증)를 분리 — 뭉치면 회사주소가 공장주소로 잘못 저장됨
            if not agg["address"] and f.get("address") and d.get("doc_type") != "factory_registration":
                agg["address"] = f["address"]
            if not agg["factory_address"] and f.get("factory_address"):
                agg["factory_address"] = f["factory_address"]
            if d.get("doc_type") == "factory_registration" and not agg["factory_address"] and f.get("address"):
                agg["factory_address"] = f["address"]   # 공장등록증의 주소는 공장 주소
            for _c in ("city", "province", "country", "zip"):        # 회사 도시/도/국가/우편(NIB)
                if not agg[_c] and f.get(_c) and d.get("doc_type") != "factory_registration":
                    agg[_c] = f[_c]
            for _fc in ("factory_city", "factory_province", "factory_country", "factory_zip"):   # 공장 도시/도/국가/우편(공장등록증)
                if not agg[_fc] and f.get(_fc):
                    agg[_fc] = f[_fc]
            # 대표자 국문/영문 분리 — 한글 포함이면 국문, 아니면 영문. 주 필드는 국문 우선.
            if f.get("responsible_person"):
                _rp = str(f["responsible_person"]).strip()
                if re.search(r'[가-힣]', _rp):
                    if not agg["responsible_person_ko"]:
                        agg["responsible_person_ko"] = _rp
                elif not agg["responsible_person_en"]:
                    agg["responsible_person_en"] = _rp
            if not agg["factory_reg_no"] and f.get("factory_reg_no"):
                agg["factory_reg_no"] = f["factory_reg_no"]
            # 전화 엄격 분리(오피스↔공장 혼입 방지): 회사 전화(office_phone)=사업자등록증(NIB)에서만,
            # 공장 전화=공장등록증에서만. 해당 문서에 전화가 없으면 비워둔다(타 문서 값으로 오염 금지).
            if d.get("doc_type") == "nib_business_license" and not agg["phone"] and f.get("phone"):
                agg["phone"] = f["phone"]
            if d.get("doc_type") == "factory_registration" and not agg["factory_phone"] and (f.get("factory_phone") or f.get("phone")):
                agg["factory_phone"] = f.get("factory_phone") or f.get("phone")
            # 사업유형은 회사 속성(사업자등록증의 업태·종목)이 정본 → NIB 우선, 없을 때만 타 문서
            if d.get("doc_type") == "nib_business_license" and f.get("business_type"):
                agg["business_type"] = f["business_type"]
            elif not agg["business_type"] and f.get("business_type"):
                agg["business_type"] = f["business_type"]
            if not agg["employee_count"] and f.get("employee_count") not in (None, ""):
                agg["employee_count"] = f["employee_count"]
            if not agg["establishment_date"] and f.get("establishment_date"):
                agg["establishment_date"] = f["establishment_date"]
            if not agg["corporate_reg_no"] and f.get("corporate_reg_no"):
                agg["corporate_reg_no"] = f["corporate_reg_no"]
        # 제품/원재료명은 해당 카탈로그 문서에서만 수집(기록·타목록의 오염 방지).
        _dt = d.get("doc_type")
        # 결정적(구조 파서) 목록이 하나라도 있으면 그 필드는 결정적 값만 받는다.
        # LLM 목록은 공급자(없음·로컬·원격)마다 달라서, 섞으면 같은 서류로 배포마다
        # 다른 원재료가 잡힌다 — 심사 결과가 배포에 따라 갈리면 안 된다.
        # (실측: GPT 는 제품×원재료 매트릭스에서 원재료 21건을 뽑았고 로컬·없음은 0건이었다.)
        # 출처는 분류 결과(field_sources) 또는 문서에 저장된 fields._sources 에서 온다.
        _src = (d.get("field_sources") or (f.get("_sources") if isinstance(f, dict) else None) or {})

        def _take(key, bucket, seenset, srcset):
            """구조 파서 값만 자동 반영. LLM 단독 값은 '확인 대기'로 돌린다.

            LLM 출력은 공급자(없음·로컬·원격)마다, 같은 모델에서도 실행마다 다르다.
            그대로 신청서에 넣으면 배포에 따라 심사 대상이 달라진다 — 사람이 한 번
            보고 넣기로 했다(결정: 재현성 > 자동화 편의)."""
            vals = [v for v in (f.get(key) or []) if v and _norm_key(v)]
            if not vals:
                return
            if _src.get(key) == "structural":
                for v in vals:
                    if _norm_key(v) not in seenset:
                        seenset.add(_norm_key(v)); bucket.append(v)
            elif _src.get(key) == "structural_reject":
                return                      # 구조 파서가 '목록 아님'이라 판정한 문서
            else:
                for v in vals:
                    if _norm_key(v) not in srcset:
                        srcset.add(_norm_key(v))
                        agg["pending"].append({"kind": key, "value": v,
                                               "doc_type": _dt, "source": "llm"})

        if _dt in _PRODUCT_SRC or _dt is None:
            _take("product_names", agg["products"], _pk, _pend_pk)
        if _dt in _MATERIAL_SRC or _dt is None:
            _take("material_names", agg["materials"], _mk, _pend_mk)
        if f.get("cert_no"):
            agg["certificates"].append({"cert_no": f.get("cert_no"), "issuer": f.get("issuer"),
                                        "expiry": f.get("expiry_date")})
    # 회사명·대표자 주 필드 = 국문 우선(없으면 영문)
    agg["company_name"] = agg["company_name_ko"] or agg["company_name_en"]
    agg["responsible_person"] = agg["responsible_person_ko"] or agg["responsible_person_en"]
    # 주소 절단 보강: LLM이 주소 뒷부분(시·도·국가)을 자른 경우 도시→도→국가 순으로 이어붙여 완성.
    agg["address"] = _complete_address(agg["address"], agg["city"], agg["province"], agg["country"])
    agg["factory_address"] = _complete_address(agg["factory_address"], agg["factory_city"], agg["factory_province"], agg["factory_country"])
    # 우편번호가 문서에 없으면 주소로 조회(전세계 지오코딩) — 회사·공장 각각.
    if not agg["zip"] and agg["address"]:
        agg["zip"] = _lookup_zip(agg["address"])
    if not agg["factory_zip"] and agg["factory_address"]:
        agg["factory_zip"] = _lookup_zip(agg["factory_address"])
    return agg


def intake_zip_iter(data, limit=200):
    """ZIP → 파일별 파싱+분류를 진행하며 진행상황을 yield (스트리밍용).

    yield ("progress", {...})  파일 처리할 때마다
    yield ("result", {...})    마지막에 전체 결과(라우팅·집계)
    """
    files = extract_zip(data)
    total = len(files[:limit])
    docs = []
    yield ("start", {"total": total})
    for i, (name, raw) in enumerate(files[:limit], 1):
        _lines = []                       # 여기서 읽은 OCR을 버리지 않고 넘긴다
        text = parse_file(name, raw, ocr_sink=_lines)
        cl = classify(name, text)
        d = {"filename": os.path.basename(name), "doc_type": cl["doc_type"],
             "ocr_lines": _lines or None,
             "doc_type_ko": DOC_KO.get(cl["doc_type"], cl["doc_type"]),
             "confidence": cl.get("confidence", 0.0), "fields": cl.get("fields", {}),
             "text_len": len(text), "excerpt": text[:300],
             "content_b64": base64.b64encode(raw).decode() if len(raw) < 3_000_000 else None,
             "content_type": _ctype(name)}
        docs.append(d)
        yield ("progress", {"done": i, "total": total, "file": d["filename"],
                            "doc_type_ko": d["doc_type_ko"]})
    found = {d["doc_type"] for d in docs}
    checklist = [{"doc_type": dt, "doc_type_ko": DOC_KO[dt], "satisfied": dt in found,
                  "files": [d["filename"] for d in docs if d["doc_type"] == dt]}
                 for dt in REQUIRED_DOCS]
    missing = [c["doc_type_ko"] for c in checklist if not c["satisfied"]]
    yield ("result", {"file_count": len(files), "classified": docs, "checklist": checklist,
                      "missing": missing, "complete": len(missing) == 0,
                      "extracted": aggregate_fields(docs)})


def intake_zip(data):
    """동기 래퍼 — 진행 이벤트를 소진하고 최종 결과만 반환."""
    res = None
    for kind, payload in intake_zip_iter(data):
        if kind == "result":
            res = payload
    return res
